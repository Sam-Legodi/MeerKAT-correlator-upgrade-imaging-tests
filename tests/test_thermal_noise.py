"""Independent radiometer examples and a real casacore MS-shaped flux path."""
import csv
import json
from itertools import combinations
from pathlib import Path

import numpy as np
import pytest
from astropy.table import Table
from docx import Document

from meerkat_corr_imaging.thermal_noise import (
    radiometer_rms_ujy, sensitivity, ms_sensitivity, thermal_noise_results, SEFD_JY)


def test_independent_radiometer_example():
    # N=2, BW=1 MHz, t=1 s, SEFD=400 Jy: sqrt(4e6)=2000 => 0.2 Jy.
    assert radiometer_rms_ujy(2, 1e6, 1, 400) == pytest.approx(200_000)
    assert radiometer_rms_ujy(2, 4e6, 1, 400) == pytest.approx(100_000)
    # Natural-weighting planning examples, not the PDF's robust=-0.5 quick-look table.
    assert radiometer_rms_ujy(58, 385e6, 3600, 425) == pytest.approx(4.439552, rel=.001)
    assert radiometer_rms_ujy(54, .7*875e6, 3600, 369) == pytest.approx(3.284523, rel=1e-6)


@pytest.mark.parametrize('values', [(1,1e6,1,400),(2,0,1,400),(2,1e6,0,400),
                                  (2,1e6,1,np.nan),(2.5,1e6,1,400)])
def test_invalid_inputs(values):
    with pytest.raises(ValueError):
        radiometer_rms_ujy(*values)


def test_band_assumptions_and_missing_metadata():
    for band, sefd in SEFD_JY.items():
        result = sensitivity(dict(antenna_count=60, effective_bandwidth_hz=1e8,
                                  on_source_integration_s=100, band=band, provenance='measured metadata'))
        assert result['sefd_jy'] == sefd
        assert result['status'] == 'available'
    assert thermal_noise_results()['mfs']['test']['theoretical_rms_ujy_beam'] is None
    with pytest.raises(ValueError, match='ambiguous'):
        sensitivity(dict(antenna_count=60, effective_bandwidth_hz=1e8,
                         on_source_integration_s=100, band='S', provenance='metadata'))


@pytest.fixture
def ms(tmp_path):
    tables = pytest.importorskip('casacore.tables')
    path = tmp_path/'test.ms'
    def write(name, columns):
        desc = []
        nrow = len(next(iter(columns.values())))
        for key, array in columns.items():
            array = np.asarray(array)
            if array.ndim == 1:
                desc.append(tables.makescacoldesc(key, array[0].item()))
            else:
                desc.append(tables.makearrcoldesc(key, array.flat[0].item(), ndim=array.ndim-1))
        tb = tables.table(str(path/name) if name else str(path), tables.maketabdesc(desc), nrow=nrow, ack=False)
        for key, array in columns.items():tb.putcol(key, np.asarray(array))
        tb.close()
    pairs = list(combinations(range(3), 2))
    ant1, ant2 = zip(*(pairs*2+[(0,3),(0,0),(0,1)]))
    # Two target dumps, separated by a gap; inactive antenna, auto, calibrator excluded.
    n = len(ant1)
    flag = np.zeros((n,4,2), bool);flag[:,0]=True;flag[6]=True
    weight = np.ones((n,4,2));weight[:,1]=0
    write('', dict(ANTENNA1=ant1, ANTENNA2=ant2,
        FIELD_ID=[0]*8+[1], DATA_DESC_ID=[0]*n, SCAN_NUMBER=[1]*3+[2]*3+[1,1,3],
        TIME=[100.]*3+[200.]*3+[100.,100.,300.], INTERVAL=[10.]*n, EXPOSURE=[10.]*n,
        FLAG_ROW=[False]*n, FLAG=flag, DATA=np.ones((n,4,2),complex),
        WEIGHT=np.ones((n,2)), WEIGHT_SPECTRUM=weight))
    write('FIELD',dict(NAME=['target','calibrator']))
    write('SPECTRAL_WINDOW',dict(CHAN_FREQ=[[1.3e9,1.301e9,1.302e9,1.303e9]],CHAN_WIDTH=[[-1e6]*4]))
    write('POLARIZATION',dict(CORR_TYPE=[[9,12]]))
    write('DATA_DESCRIPTION',dict(SPECTRAL_WINDOW_ID=[0],POLARIZATION_ID=[0]))
    return path


def test_ms_flags_weights_time_selection_and_antennas(ms):
    spec = dict(ms=str(ms),field='target')
    result = ms_sensitivity(spec)
    assert result['antenna_count'] == 3
    assert result['antenna_ids'] == [0,1,2]
    assert result['on_source_integration_s'] == 20  # not 110 s; not 6 baseline dumps
    assert result['effective_bandwidth_hz'] == 2e6  # flags AND zero spectrum weights
    assert result['flag_fraction'] == pytest.approx(5/14)
    assert result['band'] == 'L'
    assert ms_sensitivity({**spec,'scans':[2]})['on_source_integration_s'] == 10
    assert ms_sensitivity({**spec,'channels':{'0':[3]}})['effective_bandwidth_hz'] == 1e6
    with pytest.raises(ValueError,match='No usable'):
        ms_sensitivity({**spec,'frequency_range_hz':[1.299e9,1.3011e9]})
    # One of two parallel hands flagged: half the exposure, not full union bandwidth.
    from casacore.tables import table
    with table(str(ms),readonly=False,ack=False) as tb:
        flags=tb.getcol('FLAG');flags[:6,2,0]=True;tb.putcol('FLAG',flags)
    assert ms_sensitivity(spec)['effective_bandwidth_hz'] == 1.5e6


def test_complete_manifest_selection(ms,tmp_path):
    path=tmp_path/'image.manifest.json'
    path.write_text(json.dumps(dict(processing_status='complete',contract=dict(
        field_id=0, scan_ids=[2], requested_channels=[dict(spw=0,channel=3)],
        tclean_parameters=dict(vis=str(ms),datacolumn='data')))))
    result=sensitivity(dict(manifest=str(path)))
    assert result['effective_bandwidth_hz']==1e6
    assert result['on_source_integration_s']==10
    assert result['theoretical_rms_ujy_beam']==pytest.approx(radiometer_rms_ujy(3,1e6,10,425))


@pytest.mark.parametrize('change,bandwidth', [('row_flag',5e6/3), ('nonfinite_data',11e6/6),
                                            ('short_exposure',11e6/6)])
def test_ms_additional_eligibility_and_exposure(ms,change,bandwidth):
    from casacore.tables import table
    with table(str(ms),readonly=False,ack=False) as tb:
        if change=='row_flag':
            tb.putcell('FLAG_ROW',0,True)
        elif change=='nonfinite_data':
            cell=tb.getcell('DATA',0);cell[2,:]=np.nan;tb.putcell('DATA',0,cell)
        else:
            tb.putcell('EXPOSURE',0,5.)
    assert ms_sensitivity(dict(ms=str(ms),field='target'))['effective_bandwidth_hz']==pytest.approx(bandwidth)


def test_cached_visibility_approximation(tmp_path):
    import pandas as pd
    from meerkat_corr_imaging.thermal_noise import summary_sensitivity
    pd.DataFrame([dict(TIME=t,ANT1=0,ANT2=1,CLASS='cross',SCAN=1,SPW=0,POL=p,FLAG_FRAC=.5)
                  for t in [100.,110.] for p in ['XX','YY']]).to_csv(tmp_path/'perrow_amp_stats.csv',index=False)
    pd.DataFrame([dict(CLASS='cross',SPW=0,CHAN=c,FREQ_HZ=3e9+c*1e6,POL=p,
                       FLAG_FRAC=.5,RFI_FREE=False)
                  for c in range(4) for p in ['XX','YY']]).to_csv(tmp_path/'rfi_free_channel_mask.csv',index=False)
    result=summary_sensitivity(dict(visibility_results=str(tmp_path),band='S4'))
    assert result['antenna_count']==2
    assert result['on_source_integration_s']==20
    assert result['effective_bandwidth_hz']==2e6
    assert result['flag_fraction']==.5
    assert 'unavailable' in result['approximation']
    result=summary_sensitivity(dict(visibility_results=str(tmp_path),band='S4',
                                    frequency_range_hz=[3e9,3.001e9]))
    assert result['effective_bandwidth_hz']==1e6
    with pytest.raises(ValueError,match='scan'):
        summary_sensitivity(dict(visibility_results=str(tmp_path),scans=[1]))


def test_flux_report_json_csv_integration(ms,tmp_path,monkeypatch):
    from meerkat_corr_imaging import flux_analysis as flux
    from meerkat_corr_imaging.uncertainty import UncertaintyConfig
    # A representative complete flux execution; PDF exporter is separately tested.
    monkeypatch.setenv('MCI_REPORT_MANIFEST',str(tmp_path/'reports.jsonl'))
    t=Table({'RA_1':[1.,2.,3.], 'DEC_1':[-30.]*3,'RA_2':[1.,2.,3.],'DEC_2':[-30.]*3,
             'Peak_flux_1':[.1,.2,.3],'Peak_flux_2':[.11,.21,.31],
             'Total_flux_1':[.1,.2,.3],'Total_flux_2':[.11,.21,.31]})
    paths=[]
    for tag in ['low','high','mfs']:
        p=tmp_path/(tag+'.fits');t.write(p);paths.append(str(p))
    result=flux.compare_fluxes_across_band_and_scans(*paths,docx_name='flux.docx',
        uncertainty_config=UncertaintyConfig(20),
        thermal_noise={'mfs':{'test':dict(ms=str(ms),field='target')}})
    numeric=json.loads(Path(result['metrics_path']).read_text())
    assert 'low_pk' in numeric['statistics']  # existing outputs retained
    noise=numeric['thermal_noise']['mfs']['test']
    assert noise['theoretical_rms_ujy_beam']==pytest.approx(radiometer_rms_ujy(3,2e6,20,425))
    assert numeric['thermal_noise']['mfs']['reference']['status']=='unavailable'
    rows=list(csv.DictReader(open(result['thermal_noise_csv'])))
    assert len(rows)==6
    assert float(rows[-1]['theoretical_rms_ujy_beam'])==pytest.approx(noise['theoretical_rms_ujy_beam'])
    doc=Document(result['report_docx'])
    assert any('Theoretical thermal' in p.text for p in doc.paragraphs)
    assert any('RMS µJy/beam' in c.text for table in doc.tables for row in table.rows for c in row.cells)


def test_flux_step_passes_unambiguous_metadata(tmp_path,monkeypatch):
    from meerkat_corr_imaging.steps import step7_flux
    from test_packaged_steps import _config
    cfg=_config(tmp_path);cfg.tests[0].name='GPU_4kS4'
    commands=[]
    monkeypatch.setattr(step7_flux,'_run',lambda cmd,**kw:commands.append(cmd))
    step7_flux.run(cfg)
    cmd=commands[0];metadata=json.loads(cmd[cmd.index('--thermal-noise-json')+1])
    assert metadata['mfs']['test']['band']=='S4'
    assert metadata['low']['test']['frequency_range_hz']==cfg.frequency_ranges.lowband_hz
    assert metadata['mfs']['test']['ms']=='test.ms'
    cfg.extra['flux']['thermal_noise']={'mfs':{'test':{'manifest':'image.manifest.json','band':'S4'}}}
    assert step7_flux._thermal_inputs(cfg,cfg.extra['flux'],infer=True)['mfs']['test']['manifest']=='image.manifest.json'
    assert step7_flux._thermal_inputs(cfg,{},infer=False)=={}
