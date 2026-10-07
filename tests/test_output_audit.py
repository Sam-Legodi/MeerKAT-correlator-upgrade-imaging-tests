from pathlib import Path

import pytest

from meerkat_corr_imaging import cli
from meerkat_corr_imaging.config import Config, PathsCfg, Target
from meerkat_corr_imaging.output_audit import RunOutputAudit, output_scopes, snapshot
from meerkat_corr_imaging.output_paths import save_report


def config(tmp_path):
    return Config(
        project_name="outputs",
        paths=PathsCfg(reports_dir=str(tmp_path / "reports"),
                       interim_dir=str(tmp_path / "interim"),
                       sky_xmatches_dir=str(tmp_path / "matches")),
        reference=Target(name="CMC1", ms_paths=[str(tmp_path / "CMC1.ms")]),
        tests=[Target(name="GPU")],
    )


def prepare_cli(monkeypatch, cfg, runners):
    monkeypatch.setenv("TMUX", "test")
    monkeypatch.setattr(cli, "load_config", lambda _: cfg)
    monkeypatch.setattr(cli, "ALL_STEPS", tuple(runners))
    monkeypatch.setattr(cli, "STEP_RUNNERS", runners)


@pytest.mark.parametrize("failure", [None, RuntimeError, SystemExit, KeyboardInterrupt])
def test_cli_final_inventory_on_success_failure_and_interrupt(tmp_path, monkeypatch, capsys, failure):
    monkeypatch.chdir(tmp_path)
    cfg = config(tmp_path)
    outdir = tmp_path / "interim" / "CMC1"
    outdir.mkdir(parents=True)
    reused = outdir / "unchanged.csv"
    rewritten = outdir / "updated.csv"
    reused.write_text("old")
    rewritten.write_text("old")
    unrelated = tmp_path / "reports" / "unrelated.docx"
    unrelated.parent.mkdir()
    unrelated.write_text("unrelated")

    def runner(cfg):
        rewritten.write_text("new")
        (outdir / "new.csv").write_text("data")
        (outdir / "worker.log").write_text("log")
        (tmp_path / "child.log").write_text("child")
        print("RUNNER FINISHED")
        if failure:
            raise failure("synthetic failure")

    prepare_cli(monkeypatch, cfg, {"vis": ("step2_visibility_qa", runner)})
    if failure:
        expected = SystemExit if failure is RuntimeError else failure
        with pytest.raises(expected):
            cli.main(["--config", "unused", "all"])
    else:
        cli.main(["--config", "unused", "all"])
    output = capsys.readouterr().out
    assert output.count("[AUDIT] Run output summary:") == 1
    assert output.index("RUNNER FINISHED") < output.index("[AUDIT] Inputs:") < output.index("[AUDIT]  Outputs:")
    assert "[AUDIT] Step: step2_visibility_qa" in output
    assert f"NEW         {outdir} (intermediate folder)" in output
    assert f"OVERWRITTEN {outdir} (intermediate folder)" in output
    assert f"UNCHANGED   {outdir} (intermediate folder)" in output
    assert str(reused) not in output and str(rewritten) not in output
    assert str(outdir / "new.csv") not in output
    assert str(outdir / "worker.log") in output
    assert str(tmp_path / "child.log") in output
    assert str(unrelated) not in output
    assert str(tmp_path / "reports" / "produced_reports.log") in output
    logs = list((tmp_path / "reports" / "pipeline_audits").glob("*.log"))
    assert len(logs) == 1 and str(logs[0]) in output
    assert "[AUDIT] Inputs:" in logs[0].read_text()


def test_step_attribution_late_pdf_and_overwritten_pdf(tmp_path, monkeypatch, capsys):
    from docx import Document
    monkeypatch.chdir(tmp_path)
    cfg = config(tmp_path)
    reports = Path(cfg.paths.reports_dir)
    reports.mkdir()
    first = reports / "first.docx"
    first.with_suffix(".pdf").write_text("previous PDF")
    second = reports / "second.docx"

    def runner(path):
        def run(cfg):
            doc = Document()
            doc.add_paragraph(path.stem)
            save_report(doc, path)
            print(f"{path.stem} completed")
        return run

    prepare_cli(monkeypatch, cfg, {
        "first": ("first_step", runner(first)),
        "second": ("second_step", runner(second)),
    })

    def pdf(path):
        result = Path(path).with_suffix(".pdf")
        result.write_text("exported PDF")
        return result

    monkeypatch.setattr(cli, "export_pdf", pdf)
    cli.main(["--config", "unused", "all"])
    output = capsys.readouterr().out.split("[AUDIT] Run output summary:")[1]
    first_text, second_text = output.split("[AUDIT] Step: second_step")
    assert f"NEW         {first}" in first_text
    assert f"OVERWRITTEN {first.with_suffix('.pdf')}" in first_text
    assert str(second) not in first_text
    assert f"NEW         {second}" in second_text
    assert f"NEW         {second.with_suffix('.pdf')}" in second_text
    assert str(first) not in second_text


def test_pdf_failure_keeps_partial_file_and_final_summary(tmp_path, monkeypatch, capsys):
    from docx import Document
    monkeypatch.chdir(tmp_path)
    cfg = config(tmp_path)
    path = Path(cfg.paths.reports_dir) / "partial.docx"

    def runner(cfg):
        path.parent.mkdir(parents=True, exist_ok=True)
        save_report(Document(), path)

    def failing_pdf(report):
        Path(report).with_suffix(".pdf").write_text("partial")
        raise RuntimeError("export failed")

    prepare_cli(monkeypatch, cfg, {"report": ("report_step", runner)})
    monkeypatch.setattr(cli, "export_pdf", failing_pdf)
    with pytest.raises(SystemExit):
        cli.main(["--config", "unused", "all"])
    output = capsys.readouterr().out
    assert "[AUDIT] Run output summary:" in output
    assert f"NEW         {path.with_suffix('.pdf')}" in output


def test_scopes_outside_reports_for_source_matching_and_custom_report(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    cfg = config(tmp_path)
    image = tmp_path / "external" / "sky.fits"
    image.parent.mkdir()
    image.touch()
    cfg.reference.images = [str(image)]
    from meerkat_corr_imaging.pybdsf_srcfind import catalogue_paths, compute_base_name
    catalogue = catalogue_paths(str(image), compute_base_name(image.stem, None))[0]
    catalogue.parent.mkdir(parents=True)
    catalogue.touch()
    assert catalogue in snapshot(output_scopes(cfg, "src"))[0]
    match = tmp_path / "external" / "Sky-CrossMatches" / "match.fits"
    match.parent.mkdir()
    match.touch()
    cfg.extra["xmatch_pairs"] = [["a.fits", "b.fits", str(match.parent.parent / match.name)]]
    assert match in snapshot(output_scopes(cfg, "xm"))[0]
    cfg.extra["verification_report"] = {"output_docx": str(tmp_path / "custom.docx")}
    custom = tmp_path / "draft_custom_GPU.docx"
    custom.touch()
    assert custom in snapshot(output_scopes(cfg, "report"))[0]
    cfg.extra["flux"] = {"ref_low_xmatch": str(match),
                         "docx_name": str(tmp_path / "flux.docx")}
    flux_docx = tmp_path / "draft_flux.docx"
    flux_docx.touch()
    assert flux_docx in snapshot(output_scopes(cfg, "flux"))[0]


def test_invalid_survey_job_preserves_legacy_match_inventory(tmp_path, monkeypatch):
    cfg = config(tmp_path)
    match = tmp_path / "external" / "Sky-CrossMatches" / "match.fits"
    match.parent.mkdir(parents=True)
    match.touch()
    cfg.extra["xmatch_pairs"] = [["a.fits", "b.fits", str(match)]]

    def invalid_jobs(cfg):
        raise ValueError("invalid survey configuration")

    monkeypatch.setattr("meerkat_corr_imaging.survey_xmatch.resolve_survey_jobs", invalid_jobs)
    assert match in snapshot(output_scopes(cfg, "xm"))[0]


def test_same_bytes_rewrite_counts_as_overwritten_and_skipped_as_unchanged(tmp_path, capsys):
    cfg = config(tmp_path)
    outdir = tmp_path / "matches" / "crossmatched-positions"
    outdir.mkdir(parents=True)
    path = outdir / "result.json"
    path.write_text("identical")
    cfg.extra["positions"] = [{"xmatch_table": str(outdir.parent / "input.fits")}]
    with RunOutputAudit(cfg) as audit:
        with audit.step("pos", "rewrite", tmp_path / "manifest"):
            path.write_text("identical")
        with audit.step("pos", "reuse", tmp_path / "manifest"):
            pass
    output = capsys.readouterr().out
    assert f"OVERWRITTEN {path}" in output
    assert f"UNCHANGED   {path}" in output


def test_casa_intermediates_collapse_to_parent_and_keep_logs(tmp_path, capsys):
    cfg = config(tmp_path)
    cfg.casa.paired_astrometry = {"enabled": True, "output_dir": str(tmp_path / "paired")}
    outdir = tmp_path / "paired"
    with RunOutputAudit(cfg) as audit:
        with audit.step("cal", "calibration", tmp_path / "manifest"):
            table = outdir / "sky.image"
            table.mkdir(parents=True)
            (table / "table.dat").write_text("intermediate")
            (table / "table.f0").write_text("intermediate")
            (outdir / "sky.fits").write_text("deliverable")
            (table / "worker.log").write_text("log")
    output = capsys.readouterr().out
    assert f"NEW         {outdir} (intermediate folder)" in output
    assert str(table / "table.dat") not in output
    assert str(outdir / "sky.fits") in output
    assert str(table / "worker.log") in output


def test_reports_inside_interim_and_repeated_report_pdf_attribution(tmp_path, monkeypatch, capsys):
    from docx import Document
    monkeypatch.chdir(tmp_path)
    cfg = config(tmp_path)
    path = Path(cfg.paths.interim_dir) / "CMC1" / "repeated.docx"
    path.parent.mkdir(parents=True)

    def runner(cfg):
        save_report(Document(), path)

    def export_pdf(report):
        pdf = Path(report).with_suffix(".pdf")
        pdf.write_bytes(b"PDF")
        return pdf

    prepare_cli(monkeypatch, cfg, {
        "vis": ("first_step", runner),
        "vis_again": ("second_step", runner),
    })
    monkeypatch.setattr(cli, "export_pdf", export_pdf)
    # The second invocation uses the same visibility output locations.
    original_scopes = output_scopes
    monkeypatch.setattr("meerkat_corr_imaging.output_audit.output_scopes",
                        lambda cfg, command: original_scopes(cfg, "vis"))
    cli.main(["--config", "unused", "all"])
    output = capsys.readouterr().out.split("[AUDIT] Run output summary:")[1]
    first_text, second_text = output.split("[AUDIT] Step: second_step")
    assert f"NEW         {path}" in first_text
    assert str(path.with_suffix(".pdf")) not in first_text
    assert f"OVERWRITTEN {path}" in second_text
    assert f"NEW         {path.with_suffix('.pdf')}" in second_text


def test_finalization_failure_still_lists_outputs(tmp_path, monkeypatch, capsys):
    monkeypatch.chdir(tmp_path)
    cfg = config(tmp_path)
    result = Path(cfg.paths.reports_dir) / "result.json"

    def runner(cfg):
        result.write_text("result")

    prepare_cli(monkeypatch, cfg, {"vis": ("vis_step", runner)})
    original_write = Path.write_text

    def write_text(path, *args, **kwargs):
        if path.name == "produced_reports.log":
            raise OSError("cannot write report log")
        return original_write(path, *args, **kwargs)

    monkeypatch.setattr(Path, "write_text", write_text)
    with pytest.raises(OSError, match="cannot write report log"):
        cli.main(["--config", "unused", "all"])
    output = capsys.readouterr().out
    assert output.count("[AUDIT] Run output summary:") == 1
    assert f"NEW         {result}" in output
    assert "[AUDIT] Step: run_finalization" in output
    assert "[AUDIT] Step exception: OSError: cannot write report log" in output
