import json
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock

import pytest

from meerkat_corr_imaging import cli, output_paths


def test_epoch_suffix_is_idempotent():
    path = Path('/reports/interim/1785652878_sdp_l0.full.8kS4/vis_amp_summary.docx')
    result = output_paths.draft_docx_path(path)
    assert result == str(path.with_name('draft_vis_amp_summary_1785652878_sdp_l0.full.8kS4.docx'))
    assert output_paths.draft_docx_path(result) == result


@pytest.mark.parametrize('variable', ['TMUX', 'STY'])
def test_session_skips_prompt(monkeypatch, variable):
    monkeypatch.setenv(variable, 'session')
    monkeypatch.setattr('builtins.input', lambda _: pytest.fail('unexpected prompt'))
    cli._confirm_session(SimpleNamespace(config='s4.yaml', cmd='vis'), [])


@pytest.mark.parametrize('answer,continues', [('y', True), ('yes', True), ('', False), ('n', False)])
def test_session_confirmation(monkeypatch, capsys, answer, continues):
    monkeypatch.delenv('TMUX', raising=False)
    monkeypatch.delenv('STY', raising=False)
    monkeypatch.setattr('builtins.input', lambda _: answer)
    args = SimpleNamespace(config='a long config.yaml', cmd='vis')
    if continues:
        cli._confirm_session(args, ['vis', '--config', args.config])
    else:
        with pytest.raises(SystemExit):
            cli._confirm_session(args, [])
    output = capsys.readouterr().out
    session = output.split('tmux new -s ')[1].splitlines()[0]
    assert len(session) <= 10
    assert f'screen -S {session}' in output


def test_final_reports_after_audit(monkeypatch, tmp_path, capsys):
    monkeypatch.setenv('TMUX', 'session')
    cfg = SimpleNamespace(paths=SimpleNamespace(reports_dir=tmp_path), tests=[], reference=SimpleNamespace(name='epoch'))
    monkeypatch.setattr(cli, 'load_config', lambda _: cfg)
    docx = tmp_path / 'draft_epoch.docx'
    def runner(cfg):
        document = Mock()
        document.save.side_effect = lambda path: Path(path).write_bytes(b'docx')
        output_paths.save_report(document, docx)
    monkeypatch.setattr(cli, 'STEP_RUNNERS', {'vis': ('vis', runner)})
    monkeypatch.setattr(cli, 'export_pdf', lambda path: Path(path).with_suffix('.pdf'))
    cli.main(['--config', 'config.yaml', 'vis'])
    output = capsys.readouterr().out
    assert output.endswith(f'{docx}\n{docx.with_suffix(".pdf")}\n')
    assert (tmp_path / 'produced_reports.log').read_text() in output


def test_native_pdf_export_without_external_commands(monkeypatch, tmp_path):
    import subprocess
    from docx import Document
    monkeypatch.setattr(subprocess, 'Popen', lambda *a, **kw: pytest.fail('external process'))
    document = Document()
    document.add_heading('Native report', 0)
    document.add_paragraph('Epoch 1785652878, flux density and QA results.')
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = 'Metric'
    table.rows[0].cells[1].text = 'Result'
    for i in range(100):
        cells = table.add_row().cells
        cells[0].text = 'Long observation path /' + 'epoch_' * 25
        cells[1].text = str(i)
    import numpy as np
    from matplotlib import image as mpimage
    figure = tmp_path / 'diagnostic.png'
    mpimage.imsave(figure, np.zeros((20, 40, 3)))
    document.add_picture(str(figure))
    source = tmp_path / 'draft_native_epoch.docx'
    document.save(source)
    target = output_paths.export_pdf(source)
    assert target == source.with_suffix('.pdf')
    assert target.read_bytes().startswith(b'%PDF-')
    assert b'/Subtype /Image' in target.read_bytes()
    assert not list(tmp_path.glob('.mci-pdf-*'))


def test_failed_render_preserves_previous_pdf(monkeypatch, tmp_path):
    from docx import Document
    from meerkat_corr_imaging import report_pdf
    source = tmp_path / 'draft_epoch.docx'
    Document().save(source)
    target = source.with_suffix('.pdf')
    target.write_bytes(b'previous PDF')
    def fail(*args):
        raise RuntimeError('render failure')
    monkeypatch.setattr(report_pdf._Pages, 'finish_page', fail)
    with pytest.raises(RuntimeError, match='render failure'):
        output_paths.export_pdf(source)
    assert target.read_bytes() == b'previous PDF'
    assert not list(tmp_path.glob('.mci-pdf-*'))


def test_pdf_preserves_source_plot_pixels(tmp_path):
    import re
    import numpy as np
    from PIL import Image
    from docx import Document
    from docx.shared import Inches
    # Fine pixel detail must survive even when the image is scaled to fit a page.
    height, width = 900, 2400
    pixels = np.random.default_rng(42).integers(0, 256, (height, width, 3), dtype=np.uint8)
    figure = tmp_path / 'source.png'
    Image.fromarray(pixels).save(figure)
    document = Document()
    document.add_picture(str(figure), width=Inches(6.5))
    source = tmp_path / 'report.docx'
    document.save(source)
    pdf = output_paths.export_pdf(source).read_bytes()
    image_headers = re.findall(rb'/Subtype /Image\b(.*?)stream', pdf, re.S)
    dimensions = [(int(re.search(rb'/Width (\d+)', h)[1]),
                   int(re.search(rb'/Height (\d+)', h)[1])) for h in image_headers]
    assert (width, height) in dimensions
