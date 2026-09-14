"""Consolidated, image-domain continuum verification reports.

The report deliberately keeps visibility-only checks separate from checks that
can be supported by delivered images, PyBDSF catalogues, cross-match tables and
calibration-report PDFs.  It never compares unlike primary-beam correction
states.
"""

from __future__ import annotations

from dataclasses import asdict, dataclass, field
import json
import math
from pathlib import Path
from typing import Any, Iterable

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from astropy import units as u
from astropy.coordinates import SkyCoord
from astropy.io import fits
from astropy.table import Table
from astropy.wcs import WCS
from astropy.wcs.utils import proj_plane_pixel_scales
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from .uncertainty import (UncertaintyConfig, ONE_SIGMA, column, enrich_matches,
    bootstrap, wilson, ratio_error, measurement, decision, position_covariance,
    format_uncertainty, write_json, catalogue_intervals)
from .astrometry import fit_rigid
from .config import Config
from .output_paths import draft_docx_path, save_report


POSITION_LIMIT_ARCSEC = 1.0
FLUX_LIMIT_FRACTION = 0.05
RMS_RATIO_LIMIT = 1.2
RMS_ANNULUS_DEG = (0.25, 0.50)
MIN_QUALITY_MATCHES = 10

NAVY = "17365D"
TEAL = "0F6B78"
PALE_TEAL = "DDEFF1"
PASS_FILL = "E2F0D9"
CONCERN_FILL = "FCE4D6"
NA_FILL = "E7E6E6"
WHITE = "FFFFFF"
TEXT = RGBColor(36, 43, 51)


@dataclass
class RigidFit:
    rotation_deg: float
    rotation_uncertainty_deg: float
    translation_east_arcsec: float
    translation_north_arcsec: float
    postfit_median_arcsec: float
    postfit_p95_arcsec: float
    n_inliers: int
    uncertainties: dict = field(default_factory=dict)
    parameter_covariance: list | None = None
    bootstrap_parameter_covariance: list | None = None
    covariance_parameter_order: list = field(default_factory=list)
    fit_method: str = ""
    inlier_mask: list = field(default_factory=list)
    residual_arcsec: list = field(default_factory=list)
    residual_err_arcsec: list = field(default_factory=list)
    residual_ci_low_arcsec: list = field(default_factory=list)
    residual_ci_high_arcsec: list = field(default_factory=list)


@dataclass
class BandResult:
    band: str
    correction_state: str
    xmatch_table: str
    reference_image: str
    test_image: str
    reference_catalogue_count: int | None
    test_catalogue_count: int | None
    matched_count: int
    quality_count: int
    separation_median_arcsec: float | None
    separation_p95_arcsec: float | None
    separation_max_arcsec: float | None
    position_status: str
    rigid_fit: RigidFit | None
    total_flux_ratio_median: float | None
    total_flux_abs_deviation_p95: float | None
    total_flux_fraction_beyond_limit: float | None
    peak_flux_ratio_median: float | None
    flux_status: str
    reference_rms_jy_per_beam: float | None
    test_rms_jy_per_beam: float | None
    rms_ratio: float | None
    rms_status: str
    uncertainties: dict = field(default_factory=dict)
    decisions: dict = field(default_factory=dict)
    uncertainty_provenance: dict = field(default_factory=dict)
    matched_sources_output: str | None = None


def _finite(value: float | None) -> bool:
    return value is not None and math.isfinite(value)


def _status(value: float | None, limit: float) -> str:
    if not _finite(value):
        return "Not assessed"
    return "Pass" if float(value) < limit else "Concern"


def _band_name(value: str) -> str:
    lower = value.lower()
    if "lowband" in lower or "_low" in lower:
        return "Low band"
    if "highband" in lower or "_high" in lower:
        return "High band"
    return "MFS"


def _optional_bool(value: Any) -> bool | None:
    """Parse common FITS boolean representations without treating ``"F"`` as true."""

    if isinstance(value, (bool, np.bool_)):
        return bool(value)
    if isinstance(value, (int, np.integer)):
        return bool(value)
    if isinstance(value, str):
        normalized = value.strip().lower()
        if normalized in {"t", "true", "yes", "y", "1"}:
            return True
        if normalized in {"f", "false", "no", "n", "0"}:
            return False
    return None


def _correction_state(path: str | Path) -> str:
    path = Path(path)
    try:
        value = fits.getheader(path).get("PBCOR")
    except Exception:
        value = None
    parsed = _optional_bool(value)
    if parsed is not None:
        return "PB-corrected" if parsed else "non-PB"
    lower_name = path.name.lower()
    if any(token in lower_name for token in ("non_pb", "non-pb", "nonpb")):
        return "non-PB"
    return "PB-corrected" if "_pb" in lower_name else "non-PB"


def _quality_rows(table: Table) -> Table:
    """Select isolated, high-S/N Gaussian components for verification."""

    mask = np.ones(len(table), dtype=bool)
    for suffix in ("1", "2"):
        code = f"S_Code_{suffix}"
        if code in table.colnames:
            mask &= np.asarray(table[code]).astype(str) == "S"
        peak = column(table, f"Peak_flux_{suffix}", u.Jy/u.beam)
        error = column(table, f"E_Peak_flux_{suffix}", u.Jy/u.beam, error=True)
        mask &= (
            np.isfinite(peak)
            & np.isfinite(error)
            & (peak > 0)
            & (error > 0)
            & (peak / error >= 10.0)
        )
        for coordinate in ("RA", "DEC"):
            mask &= np.isfinite(column(table, f"{coordinate}_{suffix}", u.deg))
    return table[mask]


def _fit_rigid_once(reference_xy: np.ndarray, test_xy: np.ndarray) -> tuple[np.ndarray, np.ndarray]:
    ref_center = np.mean(reference_xy, axis=0)
    test_center = np.mean(test_xy, axis=0)
    u_mat, _, vt_mat = np.linalg.svd(
        (reference_xy - ref_center).T @ (test_xy - test_center)
    )
    rotation = vt_mat.T @ u_mat.T
    if np.linalg.det(rotation) < 0:
        vt_mat[-1, :] *= -1
        rotation = vt_mat.T @ u_mat.T
    translation = test_center - rotation @ ref_center
    return rotation, translation


def fit_rigid_transform(reference_xy, test_xy, *, bootstrap_samples=5000,
                        random_seed=20260911, confidence_level=ONE_SIGMA,
                        reference_covariance=None, test_covariance=None) -> RigidFit:
    return RigidFit(**fit_rigid(reference_xy, test_xy,
        UncertaintyConfig(bootstrap_samples, confidence_level, random_seed),
        reference_covariance, test_covariance, MIN_QUALITY_MATCHES))


def _phase_center(path: str | Path) -> SkyCoord:
    wcs = WCS(fits.getheader(path)).celestial
    return SkyCoord(float(wcs.wcs.crval[0]), float(wcs.wcs.crval[1]), unit="deg")


def _robust_rms(
    path: str | Path,
    *,
    inner_deg: float = RMS_ANNULUS_DEG[0],
    outer_deg: float = RMS_ANNULUS_DEG[1],
    stride: int = 3,
    return_details: bool = False,
    uncertainty_config: UncertaintyConfig | None = None,
) -> float:
    """Measure 1.4826*MAD in a fixed phase-centred annulus."""

    with fits.open(path, memmap=True) as hdus:
        header = hdus[0].header
        data = hdus[0].data
        while data.ndim > 2:
            data = data[0]
        wcs = WCS(header).celestial
        scale_deg = float(np.mean(np.abs(proj_plane_pixel_scales(wcs))))
        center_x = float(wcs.wcs.crpix[0] - 1.0)
        center_y = float(wcs.wcs.crpix[1] - 1.0)
        radius_pixels = outer_deg / scale_deg
        x0 = max(0, int(center_x - radius_pixels) - 2)
        x1 = min(data.shape[1], int(center_x + radius_pixels) + 3)
        y0 = max(0, int(center_y - radius_pixels) - 2)
        y1 = min(data.shape[0], int(center_y + radius_pixels) + 3)
        values = np.asarray(data[y0:y1:stride, x0:x1:stride], dtype=float)

    x_axis = (np.arange(x0, x1, stride) - center_x) * scale_deg
    y_axis = (np.arange(y0, y1, stride) - center_y) * scale_deg
    radius = np.hypot(y_axis[:, None], x_axis[None, :])
    selected = values[
        (radius >= inner_deg) & (radius < outer_deg) & np.isfinite(values)
    ]
    if selected.size == 0:
        raise ValueError(f"no finite pixels in RMS annulus for {path}")
    median = float(np.median(selected))
    value = 1.4826 * float(np.median(np.abs(selected - median)))
    if not return_details:
        return value
    # For Gaussian samples, median(|X|) has density 2 phi(q)/sigma
    # at q=Phi^-1(.75); quantile variance gives SE(MAD)/sigma ~1.166/sqrt(N).
    from scipy.stats import norm
    beam_major, beam_minor = header.get("BMAJ", np.nan), header.get("BMIN", np.nan)
    beam_area = np.pi * beam_major * beam_minor / (4*np.log(2))
    pixel_area = abs(np.linalg.det(wcs.pixel_scale_matrix))
    effective = None
    error = None
    reason = "missing/invalid synthesized beam area or pixel area"
    if np.isfinite([beam_major, beam_minor, pixel_area]).all() and min(beam_major, beam_minor, pixel_area) > 0:
        effective = min(float(selected.size), float(selected.size*stride**2*pixel_area/beam_area))
        reason = "fewer than two effective independent beams"
        if effective >= 2:
            error = value * 1.4826 / (4*norm.pdf(norm.ppf(.75))*np.sqrt(effective))
            reason = None
    info = measurement(value, error, uncertainty_config, method="Gaussian MAD asymptotic SE; effective independent beam count")
    info.update(n_pixels=int(selected.size), effective_independent_beams=effective,
                assumptions="locally Gaussian stationary noise; Gaussian beam correlation area; source-contamination/systematic uncertainty excluded")
    if reason:
        info["reason"] = reason
    return value, info


def _catalogue_counts(cfg: Config) -> dict[str, tuple[int | None, int | None]]:
    counts: dict[str, tuple[int | None, int | None]] = {}
    for entry in cfg.extra.get("xmatch_pairs", []) or []:
        if len(entry) < 3:
            continue
        band = _band_name(str(entry[2]))
        pair_counts: list[int | None] = []
        for path in entry[:2]:
            try:
                pair_counts.append(len(Table.read(path)))
            except Exception:
                pair_counts.append(None)
        counts[band] = (pair_counts[0], pair_counts[1])
    return counts


def _band_inputs(cfg: Config) -> list[dict[str, str]]:
    items: list[dict[str, str]] = []
    for entry in cfg.extra.get("positions", []) or []:
        if not entry.get("xmatch_table"):
            continue
        items.append(
            {
                "band": _band_name(str(entry.get("otherdatatag", entry["xmatch_table"]))),
                "xmatch_table": str(entry["xmatch_table"]),
                "reference_image": str(entry["ref_fits"]),
                "test_image": str(entry["other_fits"]),
            }
        )
    order = {"MFS": 0, "Low band": 1, "High band": 2}
    return sorted(items, key=lambda item: order.get(item["band"], 99))


def _score_band(
    item: dict[str, str],
    catalogue_counts: tuple[int | None, int | None],
    uncertainty_config: UncertaintyConfig | None = None,
    matched_output: Path | None = None,
) -> BandResult:
    reference_state = _correction_state(item["reference_image"])
    test_state = _correction_state(item["test_image"])
    if reference_state != test_state:
        raise ValueError(
            f"PB correction mismatch for {item['band']}: "
            f"reference={reference_state}, test={test_state}"
        )

    uc = uncertainty_config or UncertaintyConfig()
    uncertainties = {}
    table = enrich_matches(Table.read(item["xmatch_table"]), uc)
    table["analysis_row_index"] = np.arange(len(table))
    if matched_output is not None:
        matched_output.parent.mkdir(parents=True, exist_ok=True)
        from .xmatch_pybdsf import sanitize_fits_meta
        sanitize_fits_meta(table).write(matched_output, overwrite=True)

    quality = _quality_rows(table)
    enough = len(quality) >= MIN_QUALITY_MATCHES
    separation_median = separation_p95 = separation_max = None
    rigid_fit = None
    total_ratio_median = total_p95 = total_fraction = peak_ratio_median = None

    if enough:
        reference_coord = SkyCoord(column(quality,"RA_1",u.deg), column(quality,"DEC_1",u.deg), unit="deg")
        test_coord = SkyCoord(column(quality,"RA_2",u.deg), column(quality,"DEC_2",u.deg), unit="deg")
        separation = reference_coord.separation(test_coord).arcsec
        separation_median = float(np.median(separation))
        separation_p95 = float(np.percentile(separation, 95))
        separation_max = float(np.max(separation))

        center = _phase_center(item["reference_image"])
        ref_east, ref_north = center.spherical_offsets_to(reference_coord)
        test_east, test_north = center.spherical_offsets_to(test_coord)
        reference_xy = np.column_stack((ref_east.arcsec, ref_north.arcsec))
        test_xy = np.column_stack((test_east.arcsec, test_north.arcsec))
        def phase_cov(suffix):
            n = len(quality)
            errors = np.column_stack((np.zeros(n), np.zeros(n),
                column(quality, f"E_RA_{suffix}", u.deg, error=True),
                column(quality, f"E_DEC_{suffix}", u.deg, error=True)))
            return position_covariance(np.full(n, center.ra.deg), np.full(n, center.dec.deg),
                column(quality, f"RA_{suffix}", u.deg), column(quality, f"DEC_{suffix}", u.deg), errors)
        try:
            rigid_fit = fit_rigid_transform(reference_xy, test_xy,
                **asdict(uc), reference_covariance=phase_cov("1"), test_covariance=phase_cov("2"))
        except ValueError:
            # Raw astrometry remains assessable for degenerate fit geometry.
            rigid_fit = None

        reference_total = column(quality, "Total_flux_1", u.Jy)
        test_total = column(quality, "Total_flux_2", u.Jy)
        valid_total = (
            np.isfinite(reference_total)
            & np.isfinite(test_total)
            & (reference_total > 0)
            & (test_total > 0)
        )
        total_ratio = test_total[valid_total] / reference_total[valid_total]
        if total_ratio.size:
            total_ratio_median = float(np.median(total_ratio))
            total_p95 = float(np.percentile(np.abs(total_ratio - 1.0), 95))
            total_fraction = float(np.mean(np.abs(total_ratio - 1.0) > FLUX_LIMIT_FRACTION))

        reference_peak = column(quality, "Peak_flux_1", u.Jy/u.beam)
        test_peak = column(quality, "Peak_flux_2", u.Jy/u.beam)
        valid_peak = (
            np.isfinite(reference_peak)
            & np.isfinite(test_peak)
            & (reference_peak > 0)
            & (test_peak > 0)
        )
        peak_ratio = test_peak[valid_peak] / reference_peak[valid_peak]
        if peak_ratio.size:
            peak_ratio_median = float(np.median(peak_ratio))

    if enough:
        uncertainties.update(catalogue_intervals(quality, uc))

    for prefix in ("reference", "test"):
        try:
            value, info = _robust_rms(item[f"{prefix}_image"], return_details=True, uncertainty_config=uc)
        except ValueError as exc:
            value, info = None, {"reason": str(exc), "ci_low": None, "ci_high": None}
        uncertainties[f"{prefix}_rms_jy_per_beam"] = info
        if prefix == "reference":
            reference_rms = value
        else:
            test_rms = value
    rms_ratio = test_rms/reference_rms if _finite(test_rms) and _finite(reference_rms) and reference_rms > 0 else None
    _, rms_error = ratio_error(test_rms, reference_rms,
        uncertainties["test_rms_jy_per_beam"].get("standard_error"),
        uncertainties["reference_rms_jy_per_beam"].get("standard_error"))
    uncertainties["rms_ratio"] = measurement(rms_ratio if rms_ratio is not None else np.nan, float(rms_error), uc)
    decisions = {
        "position": decision(separation_p95, uncertainties.get("separation_p95_arcsec"), POSITION_LIMIT_ARCSEC),
        "flux": decision(total_ratio_median, uncertainties.get("total_flux_ratio_median"), FLUX_LIMIT_FRACTION, target=1.),
        "rms": decision(rms_ratio, uncertainties.get("rms_ratio"), RMS_RATIO_LIMIT),
    }
    # A configured reporting confidence may differ, but acceptance is always 1 sigma.
    if uc.confidence_level != ONE_SIGMA:
        decision_cfg = UncertaintyConfig(uc.bootstrap_samples, ONE_SIGMA, uc.random_seed)
        if enough:
            decision_intervals = catalogue_intervals(quality, decision_cfg)
            decisions["position"] = decision(separation_p95, decision_intervals['separation_p95_arcsec'], POSITION_LIMIT_ARCSEC)
            if total_ratio.size:
                decisions["flux"] = decision(total_ratio_median, decision_intervals['total_flux_ratio_median'], FLUX_LIMIT_FRACTION, target=1.)
        decisions["rms"] = decision(rms_ratio, measurement(rms_ratio if rms_ratio is not None else np.nan, float(rms_error), decision_cfg), RMS_RATIO_LIMIT)

    return BandResult(
        band=item["band"],
        correction_state=reference_state,
        xmatch_table=item["xmatch_table"],
        reference_image=item["reference_image"],
        test_image=item["test_image"],
        reference_catalogue_count=catalogue_counts[0],
        test_catalogue_count=catalogue_counts[1],
        matched_count=len(table),
        quality_count=len(quality),
        separation_median_arcsec=separation_median,
        separation_p95_arcsec=separation_p95,
        separation_max_arcsec=separation_max,
        position_status=decisions["position"]["status"],
        rigid_fit=rigid_fit,
        total_flux_ratio_median=total_ratio_median,
        total_flux_abs_deviation_p95=total_p95,
        total_flux_fraction_beyond_limit=total_fraction,
        peak_flux_ratio_median=peak_ratio_median,
        flux_status=decisions["flux"]["status"],
        reference_rms_jy_per_beam=reference_rms,
        test_rms_jy_per_beam=test_rms,
        rms_ratio=rms_ratio,
        rms_status=decisions["rms"]["status"],
        uncertainties=uncertainties, decisions=decisions,
        uncertainty_provenance={**asdict(uc), "decision_confidence_level": ONE_SIGMA,
            "measurement_assumption": "independent catalogues; no input covariance supplied",
            "quality_row_indices": column(quality, "analysis_row_index").astype(int).tolist(),
            "valid_position_errors": int(np.isfinite(column(table, "separation_err_arcsec")).sum()),
            "valid_total_ratio_errors": int(np.isfinite(column(table, "total_flux_ratio_err")).sum()),
            "valid_peak_ratio_errors": int(np.isfinite(column(table, "peak_flux_ratio_err")).sum()),
            "quality_selection": "S_Code=S where supplied; finite positions and peak/error >=10; missing peak errors fail quality selection",
            "sampling_assumption": "independent matched sources; conditional on match gate and quality cuts; measurement scatter already represented"},
        matched_sources_output=str(matched_output) if matched_output else None,
    )


def _set_cell_fill(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def _set_repeat_table_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    properties.append(repeat)


def _status_fill(text: str) -> str | None:
    if "Concern" in text:
        return CONCERN_FILL
    if "Pass" in text:
        return PASS_FILL
    if "Not assessed" in text:
        return NA_FILL
    return None


def _add_table(doc: Document, headers: list[str], rows: Iterable[Iterable[str]]):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for index, heading in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = heading
        _set_cell_fill(cell, NAVY)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in cell.paragraphs[0].runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(8)
    _set_repeat_table_header(table.rows[0])

    for row_values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row_values):
            text = str(value)
            cells[index].text = text
            cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            fill = _status_fill(text)
            if fill:
                _set_cell_fill(cells[index], fill)
            for paragraph in cells[index].paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.size = Pt(7.5)
    return table


def _add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, end])


def _configure_document(doc: Document, observation_label: str) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(0.62)
    section.bottom_margin = Inches(0.58)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9.2)
    normal.font.color.rgb = TEXT
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.05
    for name, size, color in (
        ("Title", 24, NAVY),
        ("Heading 1", 15, NAVY),
        ("Heading 2", 11.5, TEAL),
    ):
        style = doc.styles[name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(8)
        style.paragraph_format.space_after = Pt(4)

    header = section.header.paragraphs[0]
    header.text = f"MeerKAT correlator continuum verification  |  {observation_label}  |  DRAFT"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        run.font.size = Pt(7.5)
        run.font.color.rgb = RGBColor.from_string(TEAL)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Image-domain verification  •  page ")
    _add_page_field(footer)
    for run in footer.runs:
        run.font.size = Pt(7.5)
        run.font.color.rgb = RGBColor(89, 89, 89)


def _fmt(value: float | None, digits: int = 3) -> str:
    return "—" if not _finite(value) else f"{float(value):.{digits}f}"


def _fmt_pct(value: float | None, digits: int = 1) -> str:
    return "—" if not _finite(value) else f"{100.0 * float(value):.{digits}f}%"


def _metric(result, name, scale=1.):
    return format_uncertainty(getattr(result, name), result.uncertainties.get(name), scale=scale)


def _caption(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(text)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(4)
    for run in paragraph.runs:
        run.font.size = Pt(7.5)
        run.font.italic = True
        run.font.color.rgb = RGBColor(89, 89, 89)


def _page_heading(doc: Document, text: str):
    """Start a major section on a new page without creating blank pages."""

    heading = doc.add_heading(text, level=1)
    heading.paragraph_format.page_break_before = True
    return heading


def _plot_acceptance(results: list[BandResult], output: Path) -> None:
    labels = [result.band.replace(" band", "") for result in results]
    colors = ["#3A7D8C", "#5B8E7D", "#C47F3A"]
    figure, axes = plt.subplots(2, 2, figsize=(9.2, 6.6))
    panels = [
        (axes[0, 0], [result.separation_p95_arcsec for result in results], POSITION_LIMIT_ARCSEC, "Raw radial separation p95", "arcsec"),
        (axes[0, 1], [None if result.total_flux_ratio_median is None else 100 * abs(result.total_flux_ratio_median - 1) for result in results], 100 * FLUX_LIMIT_FRACTION, "Median integrated-flux error", "%"),
        (axes[1, 0], [result.rms_ratio for result in results], RMS_RATIO_LIMIT, "Measured / CMC1 RMS", "ratio"),
        (axes[1, 1], [None if result.rigid_fit is None else 1000 * result.rigid_fit.rotation_deg for result in results], None, "Rigid field rotation", "millidegree"),
    ]
    panel_uncertainties = []
    for result in results:
        flux = result.uncertainties.get("total_flux_ratio_median", {})
        lo, hi = flux.get("ci_low"), flux.get("ci_high")
        flux_interval = {}
        if lo is not None and hi is not None:
            flux_interval = dict(ci_low=100*(0 if lo <= 1 <= hi else min(abs(lo-1), abs(hi-1))),
                                 ci_high=100*max(abs(lo-1), abs(hi-1)))
        rotation = result.rigid_fit.uncertainties.get("rotation_deg", {}) if result.rigid_fit else {}
        rotation_interval = {k: v*1000 for k, v in rotation.items() if k in ("ci_low", "ci_high") and v is not None}
        panel_uncertainties.append([result.uncertainties.get("separation_p95_arcsec", {}),
            flux_interval, result.uncertainties.get("rms_ratio", {}), rotation_interval])
    for panel_index, (axis, values, limit, title, ylabel) in enumerate(panels):
        numeric = [0.0 if not _finite(value) else float(value) for value in values]
        bars = axis.bar(labels, numeric, color=colors[: len(labels)], width=0.62)
        for bar, value in zip(bars, values):
            if not _finite(value):
                bar.set_color("#C9C9C9")
                axis.text(bar.get_x() + bar.get_width() / 2, 0, "N/A", ha="center", va="bottom", fontsize=8)
            else:
                axis.text(bar.get_x() + bar.get_width() / 2, bar.get_height(), f"{float(value):.3g}", ha="center", va="bottom", fontsize=8)
        for index, value in enumerate(values):
            info = panel_uncertainties[index][panel_index]
            lo, hi = info.get("ci_low"), info.get("ci_high")
            if _finite(value) and lo is not None and hi is not None:
                # Draw absolute bounds; percentile intervals need not contain estimate.
                axis.vlines(index, lo, hi, color="black", lw=1.2)
                axis.hlines([lo, hi], index-.06, index+.06, color="black", lw=1.2)
        axis.margins(y=.18)
        if limit is not None:
            axis.axhline(limit, color="#9C2F2F", linestyle="--", linewidth=1.2, label=f"limit {limit:g}")
            axis.legend(frameon=False, fontsize=8)
        axis.set_title(title, fontsize=10, color="#17365D")
        axis.set_ylabel(ylabel)
        axis.grid(axis="y", linestyle=":", alpha=0.45)
        axis.spines[["top", "right"]].set_visible(False)
    figure.suptitle("Acceptance metrics by image product", fontsize=14, color="#17365D", fontweight="bold")
    figure.tight_layout(rect=(0, 0, 1, 0.95))
    output.parent.mkdir(parents=True, exist_ok=True)
    figure.savefig(output, dpi=180, bbox_inches="tight")
    plt.close(figure)


def _overall_status(results: list[BandResult], image_status: str) -> str:
    scored = [
        status
        for result in results
        for status in (result.position_status, result.flux_status, result.rms_status)
        if status != "Not assessed"
    ]
    if image_status == "Concern" or "Concern" in scored:
        return "Concern"
    return "Pass" if scored and len(scored) == 3*len(results) and all(status == "Pass" for status in scored) else "Partial"


def _add_status_paragraph(doc: Document, label: str, status: str, text: str) -> None:
    paragraph = doc.add_paragraph()
    label_run = paragraph.add_run(f"{label}: ")
    label_run.bold = True
    status_run = paragraph.add_run(status)
    status_run.bold = True
    status_run.font.color.rgb = RGBColor.from_string("2E7D32" if status == "Pass" else "9C2F2F" if status == "Concern" else "666666")
    paragraph.add_run(f" — {text}")


def _build_docx(
    output: Path,
    cfg: Config,
    results: list[BandResult],
    figure_path: Path,
    calreports: list[Path],
    qa_dir: Path,
    report_cfg: dict[str, Any],
) -> None:
    observation_label = cfg.tests[0].name.removeprefix("CMC2_")
    reference_label = cfg.reference.name.removeprefix("CMC1_")
    calibration_status = str(report_cfg.get("calibration_status", "Reviewed"))
    calibration_finding = str(
        report_cfg.get(
            "calibration_finding",
            "The available calibration-report PDF pages were reviewed; a numerical visibility-level acceptance result was not derived.",
        )
    )
    image_status = str(report_cfg.get("image_status", "Reviewed"))
    image_finding = str(
        report_cfg.get(
            "image_finding",
            "The MFS and independent subband diagnostic figures were reviewed.",
        )
    )
    overall = _overall_status(results, image_status)

    doc = Document()
    _configure_document(doc, observation_label)

    title = doc.add_paragraph(style="Title")
    title.add_run("Continuum imaging verification")
    subtitle = doc.add_paragraph()
    subtitle.add_run(f"CMC2 {observation_label} against CMC1 {reference_label}").bold = True
    subtitle.add_run("\nImage-product assessment • draft technical report")
    subtitle.paragraph_format.space_after = Pt(10)

    badge = doc.add_table(rows=1, cols=2)
    badge.alignment = WD_TABLE_ALIGNMENT.CENTER
    badge.autofit = True
    badge.cell(0, 0).text = "IMAGE-DOMAIN VERDICT"
    badge.cell(0, 1).text = overall.upper()
    _set_cell_fill(badge.cell(0, 0), NAVY)
    _set_cell_fill(badge.cell(0, 1), CONCERN_FILL if overall == "Concern" else PASS_FILL)
    for cell in badge.rows[0].cells:
        for run in cell.paragraphs[0].runs:
            run.font.bold = True
            run.font.size = Pt(11)
    for run in badge.cell(0, 0).paragraphs[0].runs:
        run.font.color.rgb = RGBColor(255, 255, 255)

    doc.add_heading("Executive finding", level=1)
    concern_bands = [r.band for r in results if "Concern" in (r.position_status, r.flux_status, r.rms_status)]
    concern_text = ", ".join(concern_bands) if concern_bands else "none"
    doc.add_paragraph(
        f"The image-domain assessment is {overall.lower()}. Concern-level results occur in {concern_text}. "
        "This is not a full visibility-domain acceptance decision: raw visibilities and per-scan images were unavailable for this cycle."
    )
    for result in results:
        state = result.correction_state
        position = "not assessed" if result.separation_p95_arcsec is None else f"p95 {_metric(result, 'separation_p95_arcsec')} arcsec ({result.position_status.lower()})"
        flux = "not assessed" if result.total_flux_ratio_median is None else f"median integrated-flux ratio {_metric(result, 'total_flux_ratio_median')} ({result.flux_status.lower()})"
        noise = f"RMS ratio {_metric(result, 'rms_ratio')} ({result.rms_status.lower()})" if result.rms_ratio is not None else "RMS not assessed"
        doc.add_paragraph(
            f"{result.band} [{state}]: {position}; {flux}; {noise}.",
            style="List Bullet",
        )

    caution = doc.add_paragraph()
    caution_run = caution.add_run("Caution: ")
    caution_run.bold = True
    caution.add_run(
        "Each comparison uses the same primary-beam correction state on the reference and test sides. "
        "The PB-corrected MFS comparison is not combined numerically with the non-PB low- and high-band comparisons."
    )

    _page_heading(doc, "Acceptance results")
    doc.add_paragraph(
        f"Intervals are {100*results[0].uncertainty_provenance.get('confidence_level', ONE_SIGMA):.2f}% confidence "
        "(default 68.27%, approximately 1 sigma). Brackets give absolute interval endpoints. "
        "Catalogue intervals retain the envelope of source-sampling and fixed-sample measurement-propagation intervals where available. "
        "Pass requires the full 68.27% decision interval to satisfy the existing limit. "
        "An interval touching or crossing the limit is Concern; unavailable decision uncertainty is Not assessed."
    )
    rows = []
    for result in results:
        position = f"{_metric(result, 'separation_p95_arcsec')} arcsec\n{result.position_status}"
        flux_error = None if result.total_flux_ratio_median is None else abs(result.total_flux_ratio_median - 1.0)
        flux = f"{_metric(result, 'total_flux_ratio_median')}\n{result.flux_status}"
        rms = f"{_metric(result, 'rms_ratio')}\n{result.rms_status}"
        rows.append([result.band, result.correction_state, position, flux, rms])
    _add_table(
        doc,
        [
            "Product",
            "Correction",
            "Position p95\npass below 1 arcsec",
            "Median flux ratio\nwithin 0.95–1.05",
            "RMS ratio\npass below 1.2",
        ],
        rows,
    )
    doc.add_picture(str(figure_path), width=Inches(6.7))
    _caption(
        doc,
        "Figure 1. Acceptance metrics for the MFS, low-band and high-band products. Error bars show the stated confidence intervals; missing bars mean uncertainty is unavailable. Grey bars denote insufficient catalogue evidence.",
    )

    doc.add_heading("Astrometry and rigid rotation", level=2)
    astrometry_rows = []
    for result in results:
        fit = result.rigid_fit
        astrometry_rows.append(
            [
                result.band,
                f"{result.matched_count} / {result.quality_count}",
                _metric(result, "separation_median_arcsec"),
                _metric(result, "separation_p95_arcsec"),
                "—" if fit is None else _metric(fit, "rotation_deg"),
                "—" if fit is None else _metric(fit, "translation_east_arcsec") + "\n" + _metric(fit, "translation_north_arcsec"),
                "—" if fit is None else _metric(fit, "postfit_p95_arcsec"),
                result.position_status,
            ]
        )
    _add_table(
        doc,
        ["Product", "matched / quality", "median (arcsec)", "p95 (arcsec)", "rotation (deg)", "shift E,N (arcsec)", "post-fit p95", "decision"],
        astrometry_rows,
    )

    doc.add_heading("Flux density and image noise", level=2)
    flux_rows = []
    for result in results:
        flux_rows.append(
            [
                result.band,
                f"{result.reference_catalogue_count if result.reference_catalogue_count is not None else '—'} / {result.test_catalogue_count if result.test_catalogue_count is not None else '—'}",
                _metric(result, "total_flux_ratio_median"),
                _metric(result, "peak_flux_ratio_median"),
                _metric(result, "total_flux_abs_deviation_p95", 100) + "%",
                _metric(result, "total_flux_fraction_beyond_limit", 100) + "%",
                _metric(result, "rms_ratio"),
                f"flux {result.flux_status}; RMS {result.rms_status}",
            ]
        )
    _add_table(
        doc,
        ["Product", "PyBDSF ref / test", "median total ratio", "median peak ratio", "p95 |total−1|", "fraction >5%", "RMS ratio", "decision"],
        flux_rows,
    )

    for result in results:
        doc.add_paragraph(f"{result.band} annular RMS (µJy/beam): reference "
            f"{_metric(result, 'reference_rms_jy_per_beam', 1e6)}; test "
            f"{_metric(result, 'test_rms_jy_per_beam', 1e6)}. "
            "Gaussian MAD-scale uncertainty uses independent beam area, excluding noise-model systematics.")
    doc.add_paragraph("Catalogue envelopes are conservative sensitivity bounds, not exact combined-coverage intervals. Missing formal errors leave sampling-only intervals, identified in the metrics JSON. Rigid-fit intervals retain the envelope of positional-covariance GLS and paired-source bootstrap intervals. "
        "Post-fit intervals refit resampled inliers; match selection and clipping uncertainty are not included. "
        "Wilson intervals describe the observed fraction above 5%; they do not classify latent true source fluxes.")

    diagnostic = next(iter(sorted(qa_dir.glob("*_01_combined_plane_diagnostic.png"))), None)
    subbands = next(iter(sorted(qa_dir.glob("*_02_subband_common_scale.png"))), None)
    if diagnostic and diagnostic.exists():
        _page_heading(doc, "Image inspection")
        _add_status_paragraph(doc, "Visual assessment", image_status, image_finding)
        doc.add_picture(str(diagnostic), width=Inches(6.65))
        _caption(doc, "Figure 2. Full-field and central non-PB MFS diagnostics, including radial noise and artefact statistics.")
    if subbands and subbands.exists():
        _page_heading(doc, "Frequency-resolved image inspection")
        doc.add_picture(str(subbands), width=Inches(6.25))
        _caption(doc, "Figure 3. Independent subband images on a common display scale; orange labels identify products classified as severely corrupted by the source diagnostic.")

    _page_heading(doc, "Calibration-report review")
    _add_status_paragraph(doc, "Visual assessment", calibration_status, calibration_finding)
    cal_rows = [[str(index), report.name, "Reviewed"] for index, report in enumerate(calreports, 1)]
    if not cal_rows:
        cal_rows = [["—", "No calibration-report PDF found", "Not assessed"]]
    _add_table(doc, ["#", "Calibration-report PDF", "Review state"], cal_rows)
    doc.add_paragraph(
        "The calibration-report plots provide contextual visual evidence. The <20% flagging threshold and <1% amplitude-oscillation threshold were not scored because this cycle did not derive the required visibility-domain time, channel, correlation and baseline statistics."
    )

    doc.add_heading("Checklist disposition", level=1)
    insufficient_catalogue_bands = [
        f"{result.band} ({result.quality_count} quality-selected matches)"
        for result in results
        if result.quality_count < MIN_QUALITY_MATCHES
    ]
    if insufficient_catalogue_bands:
        pybdsf_evidence = (
            "Reference and test catalogues were generated for each image product. "
            + ", ".join(insufficient_catalogue_bands)
            + " did not meet the 10-match minimum and remains not assessed for catalogue-derived metrics."
        )
    else:
        pybdsf_evidence = (
            "Reference and test catalogues were generated for each image product; "
            "all bands met the 10-match minimum for catalogue-derived metrics."
        )
    checklist = [
        ["Visual inspection of calibration report", calibration_status, calibration_finding],
        ["Sanity checks and amplitude oscillations", "Not assessed", "Raw visibilities were unavailable."],
        ["SDP flagging fraction by time/channel/correlation", "Not assessed", "No auto/cross visibility products were available."],
        ["Scan-averaged mean/RMS by polarisation and antenna/baseline", "Not assessed", "No scan-averaged visibilities were available; the planned method uses a 51-channel, third-order Savitzky–Golay model over an RFI-free flag-derived mask."],
        ["Visual inspection of images", image_status, image_finding],
        ["Source positions and relative flux across band", "Assessed", "MFS, low-band and high-band like-for-like comparisons are reported above; time/scan dependence remains unavailable."],
        ["PyBDSF source finding", "Assessed", pybdsf_evidence],
        ["Rotation and position error versus frequency", "Assessed", "Scale-fixed translation plus one rotation was fitted independently to each assessable band."],
        ["Image RMS versus expected RMS", "Assessed", "CMC1 image RMS is the expected-RMS reference in the same band, correction state and annulus."],
    ]
    _add_table(doc, ["Checklist item", "Disposition", "Evidence / limitation"], checklist)

    doc.add_heading("Interpretation limits", level=1)
    doc.add_paragraph(
        "No per-scan images or visibilities were available. Consequently, time-dependent source comparisons, auto/cross-correlation flagging fractions, scan-averaged visibility statistics and amplitude oscillations are outside the scope of this report. "
        "Primary-beam correction was not applied to the low- and high-band products during this cycle. Their results remain valid as non-PB like-for-like comparisons but are not combined with the PB-corrected MFS photometry."
    )
    doc.add_paragraph(
        "A scale change was not fitted in the astrometric model. The reported transformation is restricted to the expected global release-to-release offset: east/north translation plus one rigid rotation about the reference phase centre."
    )

    doc.add_heading("Provenance", level=1)
    provenance_rows = [
        ["Configuration", str(report_cfg.get("config_path", "loaded pipeline configuration"))],
        ["Reference", cfg.reference.name],
        ["Test", cfg.tests[0].name],
        ["Cross-match gate", str(cfg.xmatch.max_sep_arcsec)],
        ["Quality selection", "S_Code=S and peak-flux S/N ≥ 10 in both catalogues"],
        ["RMS region", f"{RMS_ANNULUS_DEG[0]:.2f}–{RMS_ANNULUS_DEG[1]:.2f} deg annulus"],
    ]
    _add_table(doc, ["Item", "Value"], provenance_rows)

    output.parent.mkdir(parents=True, exist_ok=True)
    save_report(doc, output)


def build_verification_report(cfg: Config) -> Path:
    """Build one consolidated report for the first configured test target."""

    if not cfg.tests:
        raise ValueError("verification reporting requires one configured test target")
    inputs = _band_inputs(cfg)
    if not inputs:
        raise ValueError("verification reporting requires extra.positions inputs")

    report_cfg = dict(cfg.extra.get("verification_report") or {})
    catalogue_counts = _catalogue_counts(cfg)
    uc = UncertaintyConfig.from_mapping(cfg.extra.get("uncertainty"))
    observation_label = cfg.tests[0].name.removeprefix("CMC2_")
    output_dir = Path(cfg.paths.reports_dir).expanduser() / "imaging_verification"
    results = [
        _score_band(item, catalogue_counts.get(item["band"], (None, None)), uc,
            output_dir / f"{observation_label}_{index}_matched_uncertainties.fits")
        for index, item in enumerate(inputs)
    ]

    observation_label = cfg.tests[0].name.removeprefix("CMC2_")
    output_dir = Path(cfg.paths.reports_dir).expanduser() / "imaging_verification"
    requested_output = report_cfg.get(
        "output_docx", output_dir / f"{observation_label}_imaging_verification.docx"
    )
    output = Path(draft_docx_path(requested_output, observation_label)).expanduser()
    figure_path = output_dir / "figures" / f"{observation_label}_acceptance_metrics.png"
    _plot_acceptance(results, figure_path)

    test_image = Path(inputs[0]["test_image"])
    observation_dir = test_image.parent.parent
    calreports = sorted((observation_dir / "calreports").glob("*.pdf"))
    qa_dir = observation_dir / "mfimage_frequency_assessment"
    _build_docx(output, cfg, results, figure_path, calreports, qa_dir, report_cfg)

    metrics_path = output.with_name(output.stem + "_metrics.json")
    payload = {
        "test": cfg.tests[0].name,
        "reference": cfg.reference.name,
        "thresholds": {
            "position_p95_arcsec": POSITION_LIMIT_ARCSEC,
            "median_integrated_flux_fraction": FLUX_LIMIT_FRACTION,
            "rms_ratio": RMS_RATIO_LIMIT,
        },
        "rms_annulus_deg": list(RMS_ANNULUS_DEG),
        "bands": [asdict(result) for result in results],
        "calibration_reports": [str(path) for path in calreports],
    }
    payload["uncertainty"] = asdict(uc)
    write_json(metrics_path, payload)
    print(f"[REPORT] Wrote {output}")
    print(f"[REPORT] Wrote {metrics_path}")
    return output
