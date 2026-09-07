#!/usr/bin/env python3
"""Visibility-domain verification for CASA MeasurementSets.

The analyser separates auto- and cross-correlations, derives an RFI-free
channel mask from the MeasurementSet flags, and measures scan-averaged
amplitude spectra.  Spectral RMS and amplitude oscillation are measured after
subtracting a 51-channel, third-order Savitzky--Golay trend by default.

Acceptance limits are strict: flagging fraction < 20% and fractional amplitude
oscillation < 1% are Pass; values at or above a limit are Concern.
"""

from __future__ import annotations

import argparse
import csv
from collections import defaultdict
import logging
import os
from pathlib import Path
import sys
import time
from typing import DefaultDict, Dict, Iterable, List, Mapping, Optional, Tuple

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd

from .output_paths import draft_docx_path, save_report

try:
    from casacore.tables import table as ctable
except Exception:  # Helpers remain importable and unit-testable without casacore.
    ctable = None

try:
    from scipy.signal import savgol_filter

    HAVE_SAVGOL = True
except Exception:
    HAVE_SAVGOL = False

try:
    from docx import Document
    from docx.shared import Inches

    HAVE_DOCX = True
except Exception:
    HAVE_DOCX = False


DEFAULT_SAVGOL_WINDOW = 51
DEFAULT_SAVGOL_ORDER = 3
DEFAULT_RFI_FLAG_THRESHOLD = 0.20
FLAGGING_ACCEPTANCE_LIMIT = 0.20
OSCILLATION_ACCEPTANCE_LIMIT = 0.01
CHUNK_ROWS = 2048

logging.basicConfig(
    level=logging.INFO, format="%(asctime)s [%(levelname)s] %(message)s"
)
logger = logging.getLogger(__name__)


def ensure_screen_or_tmux() -> None:
    if os.environ.get("ALLOW_NO_TMUX", "").lower() in ("1", "y", "yes", "true"):
        return
    if "STY" in os.environ or "TMUX" in os.environ:
        return
    if os.environ.get("NONINTERACTIVE_ACK", "").lower() in ("1", "y", "yes", "true"):
        return
    try:
        reply = input(
            "[WARN] Not running in screen/tmux. Continue anyway? [y/N] "
            "(set ALLOW_NO_TMUX=1 to suppress this check): "
        ).strip().lower()
    except EOFError:
        reply = "n"
    if reply not in ("y", "yes"):
        sys.exit(1)


def _require_casacore() -> None:
    if ctable is None:
        raise RuntimeError(
            "python-casacore is required to analyse MeasurementSets. "
            "Install the project's 'radio' extra or conda-forge python-casacore."
        )


def _pick_data_column(ms) -> str:
    columns = ms.colnames()
    for column in ("CORRECTED_DATA", "DATA", "MODEL_DATA"):
        if column in columns:
            return column
    raise RuntimeError(
        "No DATA-like column found among CORRECTED_DATA/DATA/MODEL_DATA"
    )


def _get_field_id(
    ms_path: str, field_sel: Optional[str]
) -> Tuple[Optional[int], Optional[str]]:
    if field_sel is None:
        return None, None
    with ctable(f"{ms_path}::FIELD", readonly=True) as field_table:
        names = field_table.getcol("NAME")
    for field_id, name in enumerate(names):
        if name == str(field_sel):
            return field_id, str(name)
    try:
        field_id = int(field_sel)
    except (TypeError, ValueError) as exc:
        raise RuntimeError(
            f'Field "{field_sel}" not found. Available: {list(names)}'
        ) from exc
    if 0 <= field_id < len(names):
        return field_id, str(names[field_id])
    raise RuntimeError(f"Field id {field_id} out of range 0..{len(names) - 1}")


def _ddid_to_spw_map(ms_path: str) -> np.ndarray:
    with ctable(f"{ms_path}::DATA_DESCRIPTION", readonly=True) as description:
        return np.asarray(description.getcol("SPECTRAL_WINDOW_ID"), dtype=int)


def _spw_chan_freqs(ms_path: str) -> Tuple[np.ndarray, np.ndarray]:
    with ctable(f"{ms_path}::SPECTRAL_WINDOW", readonly=True) as spectral_window:
        return spectral_window.getcol("CHAN_FREQ"), np.asarray(
            spectral_window.getcol("NUM_CHAN"), dtype=int
        )


def _pol_labels(ms_path: str) -> List[str]:
    with ctable(f"{ms_path}::POLARIZATION", readonly=True) as polarisation:
        corr_types = polarisation.getcol("CORR_TYPE")
    if len(corr_types) == 0:
        return []
    lookup = {
        5: "RR",
        6: "RL",
        7: "LR",
        8: "LL",
        9: "XX",
        10: "XY",
        11: "YX",
        12: "YY",
    }
    return [lookup.get(int(value), str(int(value))) for value in corr_types[0]]


def _class_name(ant1: int, ant2: int) -> str:
    return "auto" if int(ant1) == int(ant2) else "cross"


def _baseline_name(ant1: int, ant2: int) -> str:
    low, high = sorted((int(ant1), int(ant2)))
    return f"{low}-{high}"


def _ensure_odd(value: int) -> int:
    return value if value % 2 else value + 1


def derive_rfi_free_channel_mask(
    flagged: np.ndarray,
    total: np.ndarray,
    threshold: float = DEFAULT_RFI_FLAG_THRESHOLD,
) -> np.ndarray:
    """Return channels whose aggregate flagging fraction is strictly below threshold."""
    flagged = np.asarray(flagged, dtype=float)
    total = np.asarray(total, dtype=float)
    if flagged.shape != total.shape:
        raise ValueError("flagged and total arrays must have the same shape")
    if not 0.0 <= threshold <= 1.0:
        raise ValueError("RFI flag threshold must be between 0 and 1")
    fraction = np.full(flagged.shape, np.nan, dtype=float)
    np.divide(flagged, total, out=fraction, where=total > 0)
    return (total > 0) & (fraction < threshold)


def _detrend_amp(
    chan_amp: np.ndarray,
    window: int = DEFAULT_SAVGOL_WINDOW,
    order: int = DEFAULT_SAVGOL_ORDER,
    valid_mask: Optional[np.ndarray] = None,
) -> np.ndarray:
    """Return residuals, fitting the trend only through accepted finite channels."""
    amplitude = np.asarray(chan_amp, dtype=float)
    valid = np.isfinite(amplitude)
    if valid_mask is not None:
        supplied = np.asarray(valid_mask, dtype=bool)
        if supplied.shape != amplitude.shape:
            raise ValueError("valid_mask must have the same shape as chan_amp")
        valid &= supplied

    residual = np.full(amplitude.shape, np.nan, dtype=float)
    if not np.any(valid):
        return residual

    if HAVE_SAVGOL and np.count_nonzero(valid) >= order + 2:
        indices = np.arange(amplitude.size, dtype=float)
        filled = np.interp(indices, indices[valid], amplitude[valid])
        largest_odd = amplitude.size if amplitude.size % 2 else amplitude.size - 1
        requested = _ensure_odd(max(1, int(window)))
        selected_window = min(requested, largest_odd)
        if selected_window >= max(5, order + 2):
            try:
                trend = savgol_filter(
                    filled,
                    window_length=selected_window,
                    polyorder=min(int(order), selected_window - 2),
                    mode="interp",
                )
                residual[valid] = amplitude[valid] - trend[valid]
                return residual
            except Exception:
                pass

    residual[valid] = amplitude[valid] - np.nanmedian(amplitude[valid])
    return residual


def spectrum_metrics(
    amplitude: np.ndarray,
    valid_mask: np.ndarray,
    window: int = DEFAULT_SAVGOL_WINDOW,
    order: int = DEFAULT_SAVGOL_ORDER,
) -> Dict[str, float]:
    """Measure the mean, detrended RMS and fractional oscillation of a spectrum."""
    values = np.asarray(amplitude, dtype=float)
    valid = np.asarray(valid_mask, dtype=bool) & np.isfinite(values)
    count = int(np.count_nonzero(valid))
    if count == 0:
        return {"MEAN": np.nan, "RMS": np.nan, "OSC_FRAC": np.nan, "N_VALID": 0}
    mean = float(np.nanmean(values[valid]))
    residual = _detrend_amp(values, window=window, order=order, valid_mask=valid)
    rms = float(np.nanstd(residual[valid]))
    oscillation = rms / abs(mean) if np.isfinite(mean) and mean != 0.0 else np.nan
    return {
        "MEAN": mean,
        "RMS": rms,
        "OSC_FRAC": float(oscillation),
        "N_VALID": count,
    }


def _acceptance_status(value: float, limit: float) -> str:
    if not np.isfinite(value):
        return "Not assessed"
    return "Pass" if value < limit else "Concern"


def _weighted_flag_fraction(frame: pd.DataFrame) -> float:
    if frame.empty or "FLAGGED" not in frame or "TOTAL" not in frame:
        return np.nan
    total = float(frame["TOTAL"].sum())
    return float(frame["FLAGGED"].sum()) / total if total > 0 else np.nan


def acceptance_summary(
    scan_stats: pd.DataFrame,
    flagging_by_channel: pd.DataFrame,
    flag_limit: float = FLAGGING_ACCEPTANCE_LIMIT,
    oscillation_limit: float = OSCILLATION_ACCEPTANCE_LIMIT,
) -> pd.DataFrame:
    """Summarise strict Pass/Concern decisions by auto/cross class and polarisation."""
    keys = set()
    for frame in (scan_stats, flagging_by_channel):
        if not frame.empty and {"CLASS", "POL"}.issubset(frame.columns):
            keys.update(map(tuple, frame[["CLASS", "POL"]].drop_duplicates().values))

    rows = []
    for correlation_class, polarisation in sorted(keys):
        flags = flagging_by_channel[
            (flagging_by_channel["CLASS"] == correlation_class)
            & (flagging_by_channel["POL"] == polarisation)
        ]
        spectra = scan_stats[
            (scan_stats["CLASS"] == correlation_class)
            & (scan_stats["POL"] == polarisation)
        ]
        flag_fraction = _weighted_flag_fraction(flags)
        oscillations = pd.to_numeric(
            spectra.get("OSC_FRAC", pd.Series(dtype=float)), errors="coerce"
        ).dropna()
        median_osc = float(oscillations.median()) if not oscillations.empty else np.nan
        p95_osc = (
            float(np.nanpercentile(oscillations.to_numpy(), 95))
            if not oscillations.empty
            else np.nan
        )
        max_osc = float(oscillations.max()) if not oscillations.empty else np.nan
        concern_fraction = (
            float((oscillations >= oscillation_limit).mean())
            if not oscillations.empty
            else np.nan
        )
        rows.append(
            {
                "CLASS": correlation_class,
                "POL": polarisation,
                "FLAG_FRAC": flag_fraction,
                "FLAG_PCT": 100.0 * flag_fraction,
                "FLAG_LIMIT_PCT": 100.0 * flag_limit,
                "FLAG_STATUS": _acceptance_status(flag_fraction, flag_limit),
                "OSC_MEDIAN_FRAC": median_osc,
                "OSC_P95_FRAC": p95_osc,
                "OSC_MAX_FRAC": max_osc,
                "OSC_MAX_PCT": 100.0 * max_osc,
                "OSC_LIMIT_PCT": 100.0 * oscillation_limit,
                "OSC_CONCERN_FRACTION": concern_fraction,
                "OSC_STATUS": _acceptance_status(max_osc, oscillation_limit),
                "N_SCAN_BASELINE_SPECTRA": int(len(oscillations)),
            }
        )
    return pd.DataFrame(rows)


FlagCounter = DefaultDict[Tuple, np.ndarray]


def _new_flag_counter() -> FlagCounter:
    return defaultdict(lambda: np.zeros((2, 0), dtype=np.int64))


def _add_flag_counter(
    counter: FlagCounter,
    key: Tuple,
    flagged_by_pol: np.ndarray,
    total_per_pol: int,
) -> None:
    flagged = np.asarray(flagged_by_pol, dtype=np.int64)
    if counter[key].shape[1] == 0:
        counter[key] = np.zeros((2, flagged.size), dtype=np.int64)
    counter[key][0] += flagged
    counter[key][1] += int(total_per_pol)


def _flag_counter_frame(
    counter: Mapping[Tuple, np.ndarray], key_names: Iterable[str], pol_labels: List[str]
) -> pd.DataFrame:
    rows = []
    names = list(key_names)
    for key, counts in counter.items():
        key_values = key if isinstance(key, tuple) else (key,)
        base = dict(zip(names, key_values))
        for pol_index, pol in enumerate(pol_labels):
            flagged = int(counts[0, pol_index])
            total = int(counts[1, pol_index])
            fraction = flagged / total if total else np.nan
            rows.append(
                {
                    **base,
                    "POL": pol,
                    "FLAGGED": flagged,
                    "TOTAL": total,
                    "FLAG_FRAC": fraction,
                    "FLAG_PCT": 100.0 * fraction,
                    "FLAG_STATUS": _acceptance_status(
                        fraction, FLAGGING_ACCEPTANCE_LIMIT
                    ),
                }
            )
    return pd.DataFrame(rows)


def _combined_flags(table, start: int, count: int, data_shape: Tuple[int, ...]) -> np.ndarray:
    columns = table.colnames()
    if "FLAG" in columns:
        flags = np.asarray(table.getcol("FLAG", startrow=start, nrow=count), dtype=bool)
    else:
        flags = np.zeros(data_shape, dtype=bool)
    if "FLAG_ROW" in columns:
        row_flags = np.asarray(
            table.getcol("FLAG_ROW", startrow=start, nrow=count), dtype=bool
        )
        flags |= row_flags[:, np.newaxis, np.newaxis]
    return flags


def _channel_bounds(
    num_chan: np.ndarray, chan_min: Optional[int], chan_max: Optional[int]
) -> Dict[int, Tuple[int, int]]:
    bounds = {}
    for spw, channel_count in enumerate(num_chan):
        low = 0 if chan_min is None else max(0, int(chan_min))
        high = int(channel_count) if chan_max is None else min(int(channel_count), int(chan_max))
        if low >= high:
            low, high = 0, int(channel_count)
        bounds[spw] = (low, high)
    return bounds


def _write_csv(frame: pd.DataFrame, path: Path) -> Path:
    frame.to_csv(path, index=False)
    return path


def _series_label(row: pd.Series, include_spw: bool) -> str:
    return f"{row.POL} SPW{int(row.SPW)}" if include_spw else str(row.POL)


def _plot_two_classes(
    frame: pd.DataFrame,
    x_column: str,
    y_column: str,
    title: str,
    xlabel: str,
    ylabel: str,
    output: Path,
    limit: Optional[float] = None,
) -> Path:
    fig, axes = plt.subplots(1, 2, figsize=(14, 5), constrained_layout=True)
    include_spw = "SPW" in frame.columns and frame["SPW"].nunique() > 1
    group_columns = ["POL"] + (["SPW"] if include_spw else [])
    for axis, correlation_class in zip(axes, ("auto", "cross")):
        subset = frame[frame.get("CLASS", "cross") == correlation_class]
        if subset.empty:
            axis.set_title(f"{title} ({correlation_class}; no data)")
        else:
            for _, group in subset.groupby(group_columns):
                group = group.sort_values(x_column)
                axis.plot(
                    group[x_column],
                    group[y_column],
                    marker="o",
                    markersize=3,
                    linewidth=1,
                    label=_series_label(group.iloc[0], include_spw),
                )
            axis.set_title(f"{title} ({correlation_class})")
            axis.legend(fontsize="small")
        if limit is not None:
            axis.axhline(limit, color="tab:red", linestyle="--", linewidth=1, label="limit")
        axis.set_xlabel(xlabel)
        axis.set_ylabel(ylabel)
        axis.grid(linestyle=":")
    fig.savefig(output, dpi=150)
    plt.close(fig)
    return output


def _plot_category(
    frame: pd.DataFrame,
    category: str,
    value: str,
    title: str,
    ylabel: str,
    output: Path,
    limit: Optional[float] = None,
) -> Path:
    fig, axes = plt.subplots(1, 2, figsize=(16, 6), constrained_layout=True)
    for axis, correlation_class in zip(axes, ("auto", "cross")):
        subset = frame[frame.get("CLASS", "cross") == correlation_class]
        categories = sorted(subset[category].astype(str).unique())
        x = np.arange(len(categories))
        if subset.empty:
            axis.set_title(f"{title} ({correlation_class}; no data)")
        else:
            for pol in sorted(subset.POL.unique()):
                group = subset[subset.POL == pol]
                values = [
                    group[group[category].astype(str) == item][value].median()
                    if (group[category].astype(str) == item).any()
                    else np.nan
                    for item in categories
                ]
                axis.plot(x, values, marker="o", markersize=3, linewidth=1, label=pol)
            axis.set_title(f"{title} ({correlation_class})")
            axis.legend(fontsize="small")
        if len(categories) <= 64:
            axis.set_xticks(x)
            axis.set_xticklabels(categories, rotation=90)
        if limit is not None:
            axis.axhline(limit, color="tab:red", linestyle="--", linewidth=1)
        axis.set_xlabel(category.title())
        axis.set_ylabel(ylabel)
        axis.grid(linestyle=":")
    fig.savefig(output, dpi=150)
    plt.close(fig)
    return output


def analyze_single(
    ms_path: str,
    field_sel: Optional[str] = None,
    outdir: Optional[str] = None,
    window: int = DEFAULT_SAVGOL_WINDOW,
    order: int = DEFAULT_SAVGOL_ORDER,
    chan_min: Optional[int] = None,
    chan_max: Optional[int] = None,
    rfi_flag_threshold: float = DEFAULT_RFI_FLAG_THRESHOLD,
    exact_outdir: bool = False,
) -> Dict[str, Path]:
    """Analyse one MeasurementSet and write inspectable CSV, plot and DOCX products."""
    _require_casacore()
    if window != _ensure_odd(window):
        raise ValueError("Savitzky--Golay window must be odd")
    if order < 0 or window < order + 2:
        raise ValueError("Savitzky--Golay window must be at least order + 2")
    if not 0.0 <= rfi_flag_threshold <= 1.0:
        raise ValueError("RFI flag threshold must be between 0 and 1")

    ms_path = str(ms_path)
    ms_path_object = Path(ms_path)
    if outdir is None:
        outdir = ms_path_object.parent / f"{ms_path_object.with_suffix('').name}_output"
    output_dir = Path(outdir)
    if not exact_outdir and output_dir.exists() and any(output_dir.iterdir()):
        output_dir = Path(f"{output_dir}_{time.strftime('%Y-%m-%d_%H.%M.%S')}")
    output_dir.mkdir(parents=True, exist_ok=True)
    logger.info("Output: %s", output_dir)

    pol_labels = _pol_labels(ms_path)
    if not pol_labels:
        raise RuntimeError("No correlation products found in POLARIZATION table")
    npol = len(pol_labels)
    channel_frequencies, num_chan = _spw_chan_freqs(ms_path)
    ddid_to_spw = _ddid_to_spw_map(ms_path)
    bounds = _channel_bounds(num_chan, chan_min, chan_max)

    channel_flagged = {
        cls: {
            spw: np.zeros((int(num_chan[spw]), npol), dtype=np.int64)
            for spw in range(len(num_chan))
        }
        for cls in ("auto", "cross")
    }
    channel_total = {
        cls: {
            spw: np.zeros((int(num_chan[spw]), npol), dtype=np.int64)
            for spw in range(len(num_chan))
        }
        for cls in ("auto", "cross")
    }
    time_flags = _new_flag_counter()
    scan_flags = _new_flag_counter()
    antenna_flags = _new_flag_counter()
    baseline_flags = _new_flag_counter()

    field_id, field_name = _get_field_id(ms_path, field_sel)
    with ctable(ms_path, readonly=True) as main_table:
        selected = main_table if field_id is None else main_table.query(f"FIELD_ID=={field_id}")
        logger.info("Field restriction: %s", field_name or "ALL")
        data_column = _pick_data_column(selected)
        columns = selected.colnames()
        have_scan = "SCAN_NUMBER" in columns
        have_time = "TIME" in columns
        scan_numbers = (
            np.unique(selected.getcol("SCAN_NUMBER")).astype(int)
            if have_scan
            else np.array([0], dtype=int)
        )

        # Pass 1: exact FLAG/FLAG_ROW fractions and the class-specific RFI mask.
        for start in range(0, selected.nrows(), CHUNK_ROWS):
            count = min(CHUNK_ROWS, selected.nrows() - start)
            antenna1 = selected.getcol("ANTENNA1", startrow=start, nrow=count)
            antenna2 = selected.getcol("ANTENNA2", startrow=start, nrow=count)
            ddids = selected.getcol("DATA_DESC_ID", startrow=start, nrow=count)
            scans = (
                selected.getcol("SCAN_NUMBER", startrow=start, nrow=count)
                if have_scan
                else np.zeros(count, dtype=int)
            )
            times = (
                selected.getcol("TIME", startrow=start, nrow=count)
                if have_time
                else np.arange(start, start + count, dtype=float)
            )
            if "FLAG" in columns:
                # Avoid a redundant full DATA-column read during the flag-only pass.
                flags = _combined_flags(selected, start, count, (count, 0, npol))
            else:
                sample_data = selected.getcol(data_column, startrow=start, nrow=count)
                flags = _combined_flags(selected, start, count, sample_data.shape)

            for row in range(count):
                ant1, ant2 = int(antenna1[row]), int(antenna2[row])
                correlation_class = _class_name(ant1, ant2)
                baseline = _baseline_name(ant1, ant2)
                spw = int(ddid_to_spw[int(ddids[row])])
                low, high = bounds[spw]
                row_flags = flags[row, low:high, :]
                flagged_by_pol = row_flags.sum(axis=0)
                n_selected_channels = high - low
                channel_flagged[correlation_class][spw][low:high] += row_flags
                channel_total[correlation_class][spw][low:high] += 1
                _add_flag_counter(
                    time_flags,
                    (float(times[row]), correlation_class),
                    flagged_by_pol,
                    n_selected_channels,
                )
                _add_flag_counter(
                    scan_flags,
                    (int(scans[row]), correlation_class),
                    flagged_by_pol,
                    n_selected_channels,
                )
                _add_flag_counter(
                    baseline_flags,
                    (baseline, correlation_class),
                    flagged_by_pol,
                    n_selected_channels,
                )
                for antenna in {ant1, ant2}:
                    _add_flag_counter(
                        antenna_flags,
                        (antenna, correlation_class),
                        flagged_by_pol,
                        n_selected_channels,
                    )

        rfi_masks = {
            cls: {
                spw: derive_rfi_free_channel_mask(
                    channel_flagged[cls][spw],
                    channel_total[cls][spw],
                    threshold=rfi_flag_threshold,
                )
                for spw in range(len(num_chan))
            }
            for cls in ("auto", "cross")
        }

        channel_rows = []
        for correlation_class in ("auto", "cross"):
            for spw in range(len(num_chan)):
                frequencies = np.asarray(channel_frequencies[spw]).ravel()
                for channel in range(int(num_chan[spw])):
                    for pol_index, pol in enumerate(pol_labels):
                        total = int(channel_total[correlation_class][spw][channel, pol_index])
                        if total == 0:
                            continue
                        flagged = int(
                            channel_flagged[correlation_class][spw][channel, pol_index]
                        )
                        fraction = flagged / total
                        channel_rows.append(
                            {
                                "CLASS": correlation_class,
                                "SPW": spw,
                                "CHAN": channel,
                                "FREQ_HZ": float(frequencies[channel]),
                                "POL": pol,
                                "FLAGGED": flagged,
                                "TOTAL": total,
                                "FLAG_FRAC": fraction,
                                "FLAG_PCT": 100.0 * fraction,
                                "FLAG_STATUS": _acceptance_status(
                                    fraction, FLAGGING_ACCEPTANCE_LIMIT
                                ),
                                "RFI_FREE": bool(
                                    rfi_masks[correlation_class][spw][channel, pol_index]
                                ),
                            }
                        )
        flag_by_channel = pd.DataFrame(channel_rows)

        # Pass 2: average amplitudes within scan/baseline before spectral metrics.
        perrow_path = output_dir / "perrow_amp_stats.csv"
        perrow_fields = [
            "TIME",
            "ANT1",
            "ANT2",
            "BASE",
            "CLASS",
            "SCAN",
            "DDID",
            "SPW",
            "POL",
            "MEAN",
            "RMS",
            "OSC_FRAC",
            "OSC_PCT",
            "OSC_STATUS",
            "FLAG_FRAC",
            "FLAG_STATUS",
            "RFI_FREE_CHANNELS",
            "VALID_CHANNELS",
        ]
        scan_rows = []
        with perrow_path.open("w", newline="", encoding="utf-8") as stream:
            writer = csv.DictWriter(stream, fieldnames=perrow_fields)
            writer.writeheader()
            for scan_number in scan_numbers:
                scan_table = (
                    selected.query(f"SCAN_NUMBER=={int(scan_number)}")
                    if have_scan
                    else selected
                )
                accumulators: Dict[Tuple[int, int, int, int], Dict[str, np.ndarray]] = {}
                try:
                    for start in range(0, scan_table.nrows(), CHUNK_ROWS):
                        count = min(CHUNK_ROWS, scan_table.nrows() - start)
                        data = scan_table.getcol(data_column, startrow=start, nrow=count)
                        flags = _combined_flags(scan_table, start, count, data.shape)
                        antenna1 = scan_table.getcol("ANTENNA1", startrow=start, nrow=count)
                        antenna2 = scan_table.getcol("ANTENNA2", startrow=start, nrow=count)
                        ddids = scan_table.getcol("DATA_DESC_ID", startrow=start, nrow=count)
                        times = (
                            scan_table.getcol("TIME", startrow=start, nrow=count)
                            if have_time
                            else np.arange(start, start + count, dtype=float)
                        )

                        for row in range(count):
                            ant1, ant2 = int(antenna1[row]), int(antenna2[row])
                            ddid = int(ddids[row])
                            spw = int(ddid_to_spw[ddid])
                            low, high = bounds[spw]
                            correlation_class = _class_name(ant1, ant2)
                            baseline = _baseline_name(ant1, ant2)
                            amplitude = np.abs(data[row, low:high, :])
                            row_flags = flags[row, low:high, :]
                            finite = np.isfinite(amplitude)
                            rfi_mask = rfi_masks[correlation_class][spw][low:high, :]
                            valid = (~row_flags) & finite & rfi_mask

                            key = (ant1, ant2, ddid, spw)
                            if key not in accumulators:
                                shape = (high - low, npol)
                                accumulators[key] = {
                                    "SUM": np.zeros(shape, dtype=np.float64),
                                    "COUNT": np.zeros(shape, dtype=np.uint32),
                                    "FLAGGED": np.zeros(npol, dtype=np.int64),
                                    "TOTAL": np.zeros(npol, dtype=np.int64),
                                }
                            aggregate = accumulators[key]
                            aggregate["SUM"] += np.where(
                                (~row_flags) & finite, amplitude, 0.0
                            )
                            aggregate["COUNT"] += ((~row_flags) & finite).astype(np.uint32)
                            aggregate["FLAGGED"] += row_flags.sum(axis=0)
                            aggregate["TOTAL"] += high - low

                            for pol_index, pol in enumerate(pol_labels):
                                metrics = spectrum_metrics(
                                    amplitude[:, pol_index],
                                    valid[:, pol_index],
                                    window=window,
                                    order=order,
                                )
                                flag_fraction = float(row_flags[:, pol_index].mean())
                                oscillation = metrics["OSC_FRAC"]
                                writer.writerow(
                                    {
                                        "TIME": float(times[row]),
                                        "ANT1": ant1,
                                        "ANT2": ant2,
                                        "BASE": baseline,
                                        "CLASS": correlation_class,
                                        "SCAN": int(scan_number),
                                        "DDID": ddid,
                                        "SPW": spw,
                                        "POL": pol,
                                        "MEAN": metrics["MEAN"],
                                        "RMS": metrics["RMS"],
                                        "OSC_FRAC": oscillation,
                                        "OSC_PCT": 100.0 * oscillation,
                                        "OSC_STATUS": _acceptance_status(
                                            oscillation, OSCILLATION_ACCEPTANCE_LIMIT
                                        ),
                                        "FLAG_FRAC": flag_fraction,
                                        "FLAG_STATUS": _acceptance_status(
                                            flag_fraction, FLAGGING_ACCEPTANCE_LIMIT
                                        ),
                                        "RFI_FREE_CHANNELS": int(rfi_mask[:, pol_index].sum()),
                                        "VALID_CHANNELS": metrics["N_VALID"],
                                    }
                                )
                finally:
                    if have_scan:
                        scan_table.close()

                for (ant1, ant2, ddid, spw), aggregate in accumulators.items():
                    correlation_class = _class_name(ant1, ant2)
                    baseline = _baseline_name(ant1, ant2)
                    average = np.full(aggregate["SUM"].shape, np.nan, dtype=float)
                    np.divide(
                        aggregate["SUM"],
                        aggregate["COUNT"],
                        out=average,
                        where=aggregate["COUNT"] > 0,
                    )
                    low, high = bounds[spw]
                    rfi_mask = rfi_masks[correlation_class][spw][low:high, :]
                    for pol_index, pol in enumerate(pol_labels):
                        valid = (aggregate["COUNT"][:, pol_index] > 0) & rfi_mask[:, pol_index]
                        metrics = spectrum_metrics(
                            average[:, pol_index], valid, window=window, order=order
                        )
                        total = int(aggregate["TOTAL"][pol_index])
                        flagged = int(aggregate["FLAGGED"][pol_index])
                        flag_fraction = flagged / total if total else np.nan
                        oscillation = metrics["OSC_FRAC"]
                        scan_rows.append(
                            {
                                "ANT1": ant1,
                                "ANT2": ant2,
                                "BASE": baseline,
                                "CLASS": correlation_class,
                                "SCAN": int(scan_number),
                                "DDID": ddid,
                                "SPW": spw,
                                "POL": pol,
                                "MEAN": metrics["MEAN"],
                                "RMS": metrics["RMS"],
                                "OSC_FRAC": oscillation,
                                "OSC_PCT": 100.0 * oscillation,
                                "OSC_STATUS": _acceptance_status(
                                    oscillation, OSCILLATION_ACCEPTANCE_LIMIT
                                ),
                                "FLAGGED": flagged,
                                "TOTAL": total,
                                "FLAG_FRAC": flag_fraction,
                                "FLAG_STATUS": _acceptance_status(
                                    flag_fraction, FLAGGING_ACCEPTANCE_LIMIT
                                ),
                                "RFI_FREE_CHANNELS": int(rfi_mask[:, pol_index].sum()),
                                "VALID_CHANNELS": metrics["N_VALID"],
                            }
                        )

        if field_id is not None:
            selected.close()

    scan_stats = pd.DataFrame(scan_rows)
    if scan_stats.empty:
        raise RuntimeError("No visibility rows were available after selection")

    flag_by_time = _flag_counter_frame(time_flags, ["TIME", "CLASS"], pol_labels)
    if not flag_by_time.empty:
        flag_by_time["ELAPSED_HOURS"] = (
            flag_by_time["TIME"] - flag_by_time["TIME"].min()
        ) / 3600.0
    flag_by_scan = _flag_counter_frame(scan_flags, ["SCAN", "CLASS"], pol_labels)
    flag_by_antenna = _flag_counter_frame(
        antenna_flags, ["ANT", "CLASS"], pol_labels
    )
    flag_by_baseline = _flag_counter_frame(
        baseline_flags, ["BASE", "CLASS"], pol_labels
    )

    scan_group = ["SCAN", "CLASS", "POL"]
    mean_by_scan = scan_stats.groupby(scan_group, as_index=False)["MEAN"].median()
    rms_by_scan = scan_stats.groupby(scan_group, as_index=False)["RMS"].median()
    osc_by_scan = scan_stats.groupby(scan_group, as_index=False)["OSC_FRAC"].max()
    osc_by_scan["OSC_PCT"] = 100.0 * osc_by_scan["OSC_FRAC"]
    osc_by_scan["OSC_STATUS"] = osc_by_scan["OSC_FRAC"].map(
        lambda value: _acceptance_status(value, OSCILLATION_ACCEPTANCE_LIMIT)
    )

    scan_antenna_parts = [
        scan_stats.rename(columns={"ANT1": "ANT"})[
            ["SCAN", "ANT", "CLASS", "POL", "MEAN", "RMS", "OSC_FRAC"]
        ],
        scan_stats[scan_stats.ANT2 != scan_stats.ANT1]
        .rename(columns={"ANT2": "ANT"})[
            ["SCAN", "ANT", "CLASS", "POL", "MEAN", "RMS", "OSC_FRAC"]
        ],
    ]
    scan_antenna_stats = pd.concat(scan_antenna_parts, ignore_index=True).groupby(
        ["SCAN", "ANT", "CLASS", "POL"], as_index=False
    ).median(numeric_only=True)
    antenna_stats = scan_antenna_stats.groupby(
        ["ANT", "CLASS", "POL"], as_index=False
    ).median(numeric_only=True)
    baseline_stats = scan_stats.groupby(
        ["BASE", "CLASS", "POL"], as_index=False
    )[["MEAN", "RMS", "OSC_FRAC"]].median()
    for grouped_stats in (scan_antenna_stats, antenna_stats, baseline_stats):
        grouped_stats["OSC_PCT"] = 100.0 * grouped_stats["OSC_FRAC"]
        grouped_stats["OSC_STATUS"] = grouped_stats["OSC_FRAC"].map(
            lambda value: _acceptance_status(value, OSCILLATION_ACCEPTANCE_LIMIT)
        )

    acceptance = acceptance_summary(scan_stats, flag_by_channel)
    output_paths = {
        "outdir": output_dir,
        "csv_main": perrow_path,
        "scan_stats": _write_csv(
            scan_stats, output_dir / "scan_averaged_amp_stats.csv"
        ),
        "rfi_mask": _write_csv(
            flag_by_channel[
                ["CLASS", "SPW", "CHAN", "FREQ_HZ", "POL", "FLAG_FRAC", "RFI_FREE"]
            ],
            output_dir / "rfi_free_channel_mask.csv",
        ),
        "flag_by_time": _write_csv(
            flag_by_time, output_dir / "flagging_vs_time.csv"
        ),
        "flag_by_scan": _write_csv(
            flag_by_scan, output_dir / "flagging_vs_scan.csv"
        ),
        "flag_by_channel": _write_csv(
            flag_by_channel, output_dir / "flagging_by_channel.csv"
        ),
        "flag_by_antenna": _write_csv(
            flag_by_antenna, output_dir / "flagging_by_antenna.csv"
        ),
        "flag_by_baseline": _write_csv(
            flag_by_baseline, output_dir / "flagging_by_baseline.csv"
        ),
        "mean_by_scan": _write_csv(mean_by_scan, output_dir / "mean_vs_scan.csv"),
        "rms_by_scan": _write_csv(rms_by_scan, output_dir / "rms_vs_scan.csv"),
        "osc_by_scan": _write_csv(
            osc_by_scan, output_dir / "oscillation_vs_scan.csv"
        ),
        "antenna_stats": _write_csv(
            antenna_stats, output_dir / "amplitude_by_antenna.csv"
        ),
        "scan_antenna_stats": _write_csv(
            scan_antenna_stats,
            output_dir / "scan_averaged_amp_by_antenna.csv",
        ),
        "scan_baseline_stats": _write_csv(
            scan_stats[
                [
                    "SCAN",
                    "BASE",
                    "CLASS",
                    "SPW",
                    "POL",
                    "MEAN",
                    "RMS",
                    "OSC_FRAC",
                    "OSC_PCT",
                    "OSC_STATUS",
                    "FLAG_FRAC",
                    "FLAG_STATUS",
                ]
            ],
            output_dir / "scan_averaged_amp_by_baseline.csv",
        ),
        "baseline_stats": _write_csv(
            baseline_stats, output_dir / "amplitude_by_baseline.csv"
        ),
        "acceptance": _write_csv(
            acceptance, output_dir / "acceptance_summary.csv"
        ),
    }

    figures = [
        _plot_two_classes(
            mean_by_scan,
            "SCAN",
            "MEAN",
            "Scan-averaged mean amplitude",
            "Scan",
            "Median mean |V|",
            output_dir / "mean_vs_scan_split.png",
        ),
        _plot_two_classes(
            rms_by_scan,
            "SCAN",
            "RMS",
            "Scan-averaged detrended RMS",
            "Scan",
            "Median RMS",
            output_dir / "rms_vs_scan_split.png",
        ),
        _plot_two_classes(
            osc_by_scan,
            "SCAN",
            "OSC_FRAC",
            "Amplitude oscillation",
            "Scan",
            "Maximum RMS / mean",
            output_dir / "oscillation_vs_scan_split.png",
            limit=OSCILLATION_ACCEPTANCE_LIMIT,
        ),
        _plot_two_classes(
            flag_by_scan,
            "SCAN",
            "FLAG_FRAC",
            "Flagging fraction",
            "Scan",
            "Flagging fraction",
            output_dir / "flagging_vs_scan_split.png",
            limit=FLAGGING_ACCEPTANCE_LIMIT,
        ),
        _plot_two_classes(
            flag_by_time,
            "ELAPSED_HOURS",
            "FLAG_FRAC",
            "Flagging fraction",
            "Elapsed time (h)",
            "Flagging fraction",
            output_dir / "flagging_vs_time_split.png",
            limit=FLAGGING_ACCEPTANCE_LIMIT,
        ),
        _plot_two_classes(
            flag_by_channel,
            "CHAN",
            "FLAG_FRAC",
            "Flagging fraction by channel",
            "Channel",
            "Flagging fraction",
            output_dir / "flagging_by_channel.png",
            limit=FLAGGING_ACCEPTANCE_LIMIT,
        ),
        _plot_category(
            antenna_stats,
            "ANT",
            "MEAN",
            "Mean amplitude by antenna",
            "Median mean |V|",
            output_dir / "mean_per_antenna_split.png",
        ),
        _plot_category(
            antenna_stats,
            "ANT",
            "RMS",
            "Detrended RMS by antenna",
            "Median RMS",
            output_dir / "rms_per_antenna_split.png",
        ),
        _plot_category(
            flag_by_antenna,
            "ANT",
            "FLAG_FRAC",
            "Flagging by antenna",
            "Flagging fraction",
            output_dir / "flagging_per_antenna_split.png",
            limit=FLAGGING_ACCEPTANCE_LIMIT,
        ),
        _plot_category(
            baseline_stats,
            "BASE",
            "MEAN",
            "Mean amplitude by baseline",
            "Median mean |V|",
            output_dir / "mean_vs_baseline_split.png",
        ),
        _plot_category(
            baseline_stats,
            "BASE",
            "RMS",
            "Detrended RMS by baseline",
            "Median RMS",
            output_dir / "rms_vs_baseline_split.png",
        ),
        _plot_category(
            flag_by_baseline,
            "BASE",
            "FLAG_FRAC",
            "Flagging by baseline",
            "Flagging fraction",
            output_dir / "flagging_vs_baseline_split.png",
            limit=FLAGGING_ACCEPTANCE_LIMIT,
        ),
    ]

    if HAVE_DOCX:
        try:
            document = Document()
            document.add_heading("Visibility-domain verification", level=1)
            document.add_paragraph(f"MeasurementSet: {ms_path}")
            if field_name:
                document.add_paragraph(f"Field: {field_name}")
            document.add_paragraph(
                "Amplitude spectra were averaged within each scan and baseline. "
                f"Channels with an aggregate flagging fraction >= "
                f"{100 * rfi_flag_threshold:.0f}% were excluded. RMS was measured "
                f"after subtracting a {window}-channel, order-{order} "
                "Savitzky–Golay trend."
            )
            document.add_paragraph(
                "Pass requires flagging <20% and fractional amplitude oscillation "
                "<1%. The oscillation decision is conservative: any assessed "
                "scan/baseline spectrum at or above 1% gives Concern for that "
                "auto/cross class and polarisation."
            )
            document.add_heading("Acceptance summary", level=2)
            table = document.add_table(rows=1, cols=6)
            headings = [
                "Class",
                "Pol",
                "Flagging",
                "Decision",
                "Maximum oscillation",
                "Decision",
            ]
            for cell, heading in zip(table.rows[0].cells, headings):
                cell.text = heading
            for row in acceptance.itertuples(index=False):
                cells = table.add_row().cells
                values = [
                    row.CLASS,
                    row.POL,
                    f"{row.FLAG_PCT:.2f}%",
                    row.FLAG_STATUS,
                    f"{row.OSC_MAX_PCT:.2f}%",
                    row.OSC_STATUS,
                ]
                for cell, value in zip(cells, values):
                    cell.text = str(value)
            document.add_heading("Diagnostic figures", level=2)
            for figure in figures:
                document.add_paragraph(figure.name)
                document.add_picture(str(figure), width=Inches(6.5))
            report_path = Path(
                draft_docx_path(output_dir / "vis_amp_summary.docx", output_dir.name)
            )
            save_report(document, report_path)
            output_paths["report"] = report_path
        except Exception as exc:
            logger.warning("Could not create Word report: %s", exc)
    else:
        logger.info("python-docx not available; skipping DOCX report")

    for index, figure in enumerate(figures):
        output_paths[f"figure_{index}"] = figure
    return output_paths


def compare_results(
    result_a: Dict[str, Path], result_b: Dict[str, Path], outdir: Path
) -> None:
    """Compare scan-averaged reference and test visibility statistics."""
    try:
        a = pd.read_csv(result_a["outdir"] / "scan_averaged_amp_stats.csv")
        b = pd.read_csv(result_b["outdir"] / "scan_averaged_amp_stats.csv")
    except Exception as exc:
        logger.warning("Comparison skipped: %s", exc)
        return

    summary_keys = ["CLASS", "POL"]
    summaries = []
    for label, frame in (("A", a), ("B", b)):
        grouped = frame.groupby(summary_keys, as_index=False)[
            ["MEAN", "RMS", "OSC_FRAC", "FLAG_FRAC"]
        ].median()
        summaries.append(
            grouped.rename(
                columns={column: f"{column}_{label}" for column in grouped.columns if column not in summary_keys}
            )
        )
    merged_summary = pd.merge(summaries[0], summaries[1], on=summary_keys, how="outer")
    for metric in ("MEAN", "RMS", "OSC_FRAC", "FLAG_FRAC"):
        merged_summary[f"{metric}_DELTA_B_MINUS_A"] = (
            merged_summary[f"{metric}_B"] - merged_summary[f"{metric}_A"]
        )
    merged_summary.to_csv(outdir / "compare_dataset_summary.csv", index=False)

    for metric in ("MEAN", "RMS", "OSC_FRAC", "FLAG_FRAC"):
        group_keys = ["SCAN", "CLASS", "POL"]
        grouped_a = a.groupby(group_keys, as_index=False)[metric].median().rename(
            columns={metric: f"{metric}_A"}
        )
        grouped_b = b.groupby(group_keys, as_index=False)[metric].median().rename(
            columns={metric: f"{metric}_B"}
        )
        pd.merge(grouped_a, grouped_b, on=group_keys, how="outer").sort_values(
            ["CLASS", "POL", "SCAN"]
        ).to_csv(outdir / f"compare_{metric.lower()}_vs_scan.csv", index=False)


def main() -> None:
    ensure_screen_or_tmux()
    parser = argparse.ArgumentParser(
        description="Analyse flagging and scan-averaged amplitude oscillations in a MeasurementSet."
    )
    parser.add_argument("--ms", required=True, help="Path to .ms or .mms")
    parser.add_argument("--ms-ref", default=None, help="Optional reference MS to compare")
    parser.add_argument("--field", default=None, help="FIELD name or integer FIELD_ID")
    parser.add_argument("--outdir", default=None, help="Output directory")
    parser.add_argument(
        "--exact-outdir",
        action="store_true",
        help=(
            "Write to the exact --outdir and replace same-named generated files; "
            "used by the pipeline for deterministic reference reuse"
        ),
    )
    parser.add_argument(
        "--window", type=int, default=DEFAULT_SAVGOL_WINDOW, help="Odd Savitzky–Golay window; default 51"
    )
    parser.add_argument(
        "--order", type=int, default=DEFAULT_SAVGOL_ORDER, help="Savitzky–Golay order; default 3"
    )
    parser.add_argument("--chan-min", type=int, default=None, help="First channel, inclusive")
    parser.add_argument("--chan-max", type=int, default=None, help="Last channel, exclusive")
    parser.add_argument(
        "--rfi-flag-threshold",
        type=float,
        default=DEFAULT_RFI_FLAG_THRESHOLD,
        help="Exclude channels with aggregate flagging at or above this fraction; default 0.20",
    )
    parser.add_argument(
        "--ms-ref-results",
        default=None,
        help="Existing reference-result directory; avoids re-reading --ms-ref",
    )
    args = parser.parse_args()

    try:
        result_test = analyze_single(
            ms_path=args.ms,
            field_sel=args.field,
            outdir=args.outdir,
            window=args.window,
            order=args.order,
            chan_min=args.chan_min,
            chan_max=args.chan_max,
            rfi_flag_threshold=args.rfi_flag_threshold,
            exact_outdir=args.exact_outdir,
        )
        if args.ms_ref_results:
            reference_outdir = Path(args.ms_ref_results)
            required = reference_outdir / "scan_averaged_amp_stats.csv"
            if not required.is_file():
                raise RuntimeError(
                    f"Reference results are incomplete; missing {required}"
                )
            result_reference = {"outdir": reference_outdir}
            compare_results(result_test, result_reference, outdir=result_test["outdir"])
        elif args.ms_ref:
            result_reference = analyze_single(
                ms_path=args.ms_ref,
                field_sel=args.field,
                window=args.window,
                order=args.order,
                chan_min=args.chan_min,
                chan_max=args.chan_max,
                rfi_flag_threshold=args.rfi_flag_threshold,
            )
            compare_results(result_test, result_reference, outdir=result_test["outdir"])
        logger.info("Done")
    except Exception as exc:
        logger.error("Error: %s", exc)
        sys.exit(2)


if __name__ == "__main__":
    main()
