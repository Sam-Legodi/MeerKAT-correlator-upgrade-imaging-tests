#!/usr/bin/env python3
"""Refresh per-source error bars in existing DRAFT_05 Section 4.2 reports.

Run with --reports-dir pointing to consolidated_reports. Existing matched rows,
figure captions, layout and centre-comparison plots are retained. The input
centre_statistics.json supplies an independent invariant check for all products.
"""
import argparse
import hashlib
import json
from pathlib import Path
import re
import zipfile

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import numpy as np
from astropy.io import fits
from astropy.table import Table
from docx import Document
from lxml import etree

from meerkat_corr_imaging.uncertainty import enrich_matches, UncertaintyConfig
from meerkat_corr_imaging.positions_analysis import _add_enclosing_ellipse
from meerkat_corr_imaging.xmatch_pybdsf import sanitize_fits_meta


PREFIXES = ('DRA-vs-phi', 'DDEC-vs-phi', 'DRA-vs-DDEC-bmaj',
            'DRA-vs-DDEC-arcsec', 'theta-vs-phi', 'rho-over-bmaj-vs-phi')
TAGS = {'MFS': 'mfs', 'Low band': 'low', 'High band': 'high'}


def summary(values):
    mean, median = float(np.mean(values)), float(np.median(values))
    sd = float(np.std(values))
    mad = float(np.median(np.abs(values-median)))
    return dict(mean=mean, median=median, sd_population=sd, mad_raw=mad,
                mean_se=sd/np.sqrt(len(values)),
                median_se_normal_approx=np.sqrt(np.pi/2)*1.482602218505602*mad/np.sqrt(len(values)))


def ellipse(ax, x, y, colour, label, unit=''):
    _add_enclosing_ellipse(ax, x, y, colour, label)
    ax.lines[-1].set_label(
        f'{label}\ncenter [{np.mean(x):.2f}, {np.mean(y):.2f}]\n'
        f'scatter SD [{np.std(x):.2f}, {np.std(y):.2f}]'+unit)


def plot_product(table, bmaj, mode, tag, index, path):
    e, n = (np.asarray(table[k]) for k in ('east_offset_arcsec', 'north_offset_arcsec'))
    ee, ne = (np.asarray(table[k]) for k in ('east_offset_err_arcsec', 'north_offset_err_arcsec'))
    phi = [np.asarray(table[f'field_angle_{s}_deg']) for s in (1, 2)]
    pe = [np.asarray(table[f'field_angle_err_{s}_deg']) for s in (1, 2)]
    fig, ax = plt.subplots(figsize=(12, 5.5))
    fig.subplots_adjust(left=.08, right=.73, bottom=.15, top=.88)
    if index in (2, 3):
        scale = bmaj if index == 2 else 1.
        x, y = e/scale, n/scale
        ax.plot(x, y, 'k.', alpha=.7, label=f'{mode} {tag.upper()} matched sources')
        ellipse(ax, x, y, 'black', 'Descriptive', ' arcsec' if index == 3 else '')
        ax.errorbar(x, y, xerr=ee/scale, yerr=ne/scale, fmt='none', alpha=.3, color='gray')
        if index == 3 and tag == 'mfs':
            ax.axvline(np.mean(x), color='red', lw=1.2, zorder=3)
            ax.axhline(np.mean(y), color='red', lw=1.2, zorder=3)
        ax.set_aspect('equal', adjustable='datalim')
        ax.set(xlabel='ΔRA / Bmaj' if index == 2 else 'ΔRA (arcsec)',
               ylabel='ΔDec / Bmaj' if index == 2 else 'ΔDec (arcsec)',
               title=f'{mode} {tag.upper()} sky residuals')
    else:
        if index == 0:
            x, xe, label = e/bmaj, ee/bmaj, 'ΔRA / Bmaj'
        elif index == 1:
            x, xe, label = n/bmaj, ne/bmaj, 'ΔDec / Bmaj'
        elif index == 4:
            x = np.asarray(table['offset_angle_deg'])/bmaj
            xe = np.asarray(table['offset_angle_err_deg'])/bmaj
            label = 'θ (deg) / Bmaj'
        else:
            x = np.asarray(table['separation_arcsec'])/bmaj
            label = 'ρ / Bmaj'
        for i, (y, ye, colour, marker) in enumerate(zip(phi, pe, ('red', 'black'), ('+', '.'))):
            ax.plot(x, y, color=colour, marker=marker, ls='none', alpha=.7, label=f'φ{i+1}')
            ellipse(ax, x, y, colour, f'φ{i+1} descriptive')
            if index == 5:
                ax.hlines(y, np.asarray(table['separation_ci_low_arcsec'])/bmaj,
                          np.asarray(table['separation_ci_high_arcsec'])/bmaj, alpha=.3, color='gray')
            else:
                ax.errorbar(x, y, xerr=xe, yerr=ye, fmt='none', alpha=.3, color='gray')
        ax.set(xlabel=label, ylabel='φ (deg)', title=f'{mode} {tag.upper()} — {label} versus φ')
    ax.grid(ls=':', alpha=.5)
    legend = ax.legend(loc='center left', bbox_to_anchor=(1.02, .5), fontsize=9)
    fig.canvas.draw()
    box = legend.get_window_extent(fig.canvas.get_renderer())
    assert box.x1 <= fig.bbox.x1 and box.y1 <= fig.bbox.y1, (mode, tag, index, box)
    fig.savefig(path, dpi=220)
    plt.close(fig)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--reports-dir', type=Path, required=True)
    args = parser.parse_args()
    root = args.reports_dir.resolve()
    out = root/'report_refresh/section42_mode_comparisons'
    centre_path = out/'centre_statistics.json'
    before_centres = centre_path.read_bytes()
    products = json.loads(before_centres)['products']
    plt.rcParams.update({'font.size': 12, 'axes.labelsize': 13, 'axes.titlesize': 14})
    audit, images = [], {}
    for product in products:
        mode_id, tag = product['mode'], product['tag']
        mode = mode_id.split('_', 1)[1]
        path = Path(product['table'])
        old = Table.read(path)
        new = enrich_matches(old, UncertaintyConfig(), ra_error_convention='on_sky')
        for name in ('east_offset_arcsec', 'north_offset_arcsec', 'separation_arcsec'):
            assert np.array_equal(old[name], new[name], equal_nan=True), (mode_id, tag, name)
        for axis, key in [('east', 'east_offset_arcsec'), ('north', 'north_offset_arcsec')]:
            actual = summary(np.asarray(new[key]))
            for name, expected in product['summary'][axis].items():
                assert np.isclose(actual[name], expected, rtol=1e-12, atol=1e-12), (mode_id, tag, axis, name)
        metrics_path = next((root/'report_refresh'/mode_id/'reports/imaging_verification').glob('*metrics.json'))
        bands = json.loads(metrics_path.read_text())['bands']
        band = next(b for b in bands if TAGS[b['band']] == tag)
        bmaj = float(fits.getheader(band['reference_image'])['BMAJ'])*3600
        for name, err in [('east_offset_bmaj', 'east_offset_err_arcsec'),
                          ('north_offset_bmaj', 'north_offset_err_arcsec'),
                          ('radial_offset_bmaj', 'separation_err_arcsec')]:
            new[name+'_err'] = np.asarray(new[err])/bmaj
        new['offset_angle_deg_err'] = new['offset_angle_err_deg']
        ratio = np.asarray(new['east_offset_err_arcsec'])/np.asarray(old['east_offset_err_arcsec'])
        audit.append(dict(mode=mode_id, product=tag, n=len(new), offsets_unchanged=True,
                          centre_statistics_unchanged=True,
                          median_east_error_factor=float(np.nanmedian(ratio)),
                          ra_error_convention='on_sky', reference_image=band['reference_image']))
        sanitize_fits_meta(new).write(path, overwrite=True)
        for index, prefix in enumerate(PREFIXES):
            figure = out/f'{mode_id}_{tag}_corrected_{index}.png'
            plot_product(new, bmaj, mode, tag, index, figure)
            images[(mode, tag, index)] = figure
            # Keep previously linked figure assets consistent with the refreshed report.
            original = path.parent/f'{prefix}-refx{mode_id}_{tag}.png'
            if original.exists():
                original.write_bytes(figure.read_bytes())
            wide = out/f'{mode_id}_{tag}_wide_{index}.png'
            if wide.exists():
                wide.write_bytes(figure.read_bytes())
        print(mode_id, tag, 'offsets and centre errors unchanged', flush=True)
    assert centre_path.read_bytes() == before_centres
    for docpath in sorted(root.glob('DRAFT_05_*astrometry.docx')):
        doc = Document(docpath)
        replacement, latest, count = {}, None, 0
        for paragraph in doc.paragraphs:
            for blip in paragraph._p.xpath('.//a:blip'):
                rid = blip.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
                latest = doc.part.related_parts[rid]
            match = re.match(r'Figure (\d+)\. (\w+), (MFS|Low band|High band):', paragraph.text)
            if match:
                number, mode, label = match.groups()
                index = (int(number)-10) % 6
                replacement[str(latest.partname).lstrip('/')] = images[(mode, TAGS[label], index)].read_bytes()
                count += 1
        assert count == 72, (docpath, count)
        # Explain corrected formal errors without changing statistical definitions.
        note = ('Catalogue error convention: PyBDSF E_RA is an on-sky angular uncertainty. '
                'The propagated source error bars use this convention without an additional cos(dec) factor. '
                'Measured offsets, source scatter and the scatter-based mean and median errors are unchanged.')
        for paragraph in doc.paragraphs:
            if paragraph.text.startswith('Input provenance:'):
                if note not in paragraph.text:
                    paragraph.add_run('\n'+note)
                break
        replacement['word/document.xml'] = etree.tostring(doc.element, xml_declaration=True, encoding='UTF-8', standalone=True)
        temporary = docpath.with_suffix('.tmp.docx')
        with zipfile.ZipFile(docpath) as oldzip, zipfile.ZipFile(temporary, 'w', compression=zipfile.ZIP_DEFLATED) as newzip:
            for item in oldzip.infolist():
                newzip.writestr(item, replacement.get(item.filename, oldzip.read(item.filename)))
        temporary.replace(docpath)
        print(docpath.name, count, 'plots refreshed', flush=True)
    payload = dict(products=audit, centre_statistics_sha256=hashlib.sha256(before_centres).hexdigest(),
                   centre_statistics_file_unchanged=True, regenerated_figures=144,
                   unaffected_figures='Combined-mode centre and angle plots contain no propagated catalogue errors')
    (out/'pybdsf_ra_error_correction_audit.json').write_text(json.dumps(payload, indent=2)+'\n')


if __name__ == '__main__':
    main()
